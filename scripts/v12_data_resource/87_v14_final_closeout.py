#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""v14 final close-out: GSVA citation, Zenodo guide refresh, three-cold-review response document,
manuscript rebuild, ZIPs and push."""
import os, re, subprocess, shutil, zipfile, csv
V14=r"C:\Users\fengq\Desktop\EGFR\EGFR的v14"; PK=os.path.join(V14,"dataset_package")
# 1. GSVA citation in the deliverable markdown
mdp=os.path.join(V14,"01_manuscript","DataDescriptor_v14.md"); md=open(mdp,encoding="utf-8").read()
if "GSVA7" not in md:
    md=re.sub(r"(GSVA) value", r"\g<1>7 value", md, count=1)
    if "GSVA7" not in md:
        md=md.replace("it is a GSVA value","it is a GSVA7 value",1)
    open(mdp,"w",encoding="utf-8").write(md)
print("GSVA cited:", "GSVA7" in open(mdp,encoding="utf-8").read())
# 2. Zenodo guide -> v14
g=os.path.join(V14,"03_Zenodo上传指南.md"); t=open(g,encoding="utf-8").read()
t=t.replace("v13 数据资源包","v14 数据资源包").replace("★Zenodo上传-就选这个_v13","★Zenodo上传-就选这个_v14")
t=t.replace("EGFR的v13","EGFR的v14").replace("DataDescriptor_v13","DataDescriptor_v14")
t=re.sub(r"26 case-control expression cohorts \(2,711 assay records, 1,086 patients with identifiers, 12 platforms[^)]*\)",
         "26 case-control expression cohorts (2,711 assay records, 1,086 patients with identifiers, 12 platforms)", t)
open(g,"w",encoding="utf-8").write(t); print("Zenodo guide updated to v14")
# 3. three-cold-review response document
doc = """# 三路独立冷读（对 v14）意见与处置

> 方法：三名上下文全新、互不知情的审稿人（Scientific Data 编辑合规 / 生物统计 / 数据工程与可复现性）冷读同一份 v14 交付物，
> 允许读文件、跑只读命令、在临时目录解压复算，禁止修改文件。三人共做：84 条 SHA-256 全量复核、ZIP↔目录逐文件比对、
> Table 1 逐格重算、两个示例实跑并与 expected_output 逐字节比对、SMDH 效应量独立重算（差 1e-16）、四张图自写脚本独立审查。
> 结论：**未发现新的 Blocker**（校验链与数字全部通过），但提出 **6 条 Major + 25 条 Minor**，多为元数据/表述/包自足性问题。
> 下面逐条给出**我的复核与处置**。

## A. 已在本轮修正（有据可查）

| # | 意见 | 我的复核 | 处置 |
|---|---|---|---|
| A1 | **包内自述版本仍是 v13**（`00_VERSION.txt`、README、CHANGELOG）| 成立 | 全部改为 **v14.0** 并补 v14 变更条目 |
| A2 | **文件数写 82，实际 86** | 成立 | 正文改为**由清单自动生成**（86 文件 / 84 条 SHA-256，并说明两条自引用）|
| A3 | **主合并表含 k=1 行却像"合并估计"**（5 行；`meta_unpaired_only` 38 行）| 成立 | 全部 k=1 行标注 **"single cohort; not a pooled estimate"**，p/FDR/CI/PI 置为 not applicable；正文写明该规则（原先"requires at least two cohorts"与表不符）|
| A4 | **per_state_evidence_matrix 仍是 v11 旧值，与主表 51 行不一致**（CRC/IBD/STAD 各 17）| **成立，属实质错误** | 用 v13 修正后的逐队列效应**全量重算**（170 行），并另存 `per_state_evidence_matrix_v14recomputed.csv` |
| A5 | **输入数据无法从稿件追溯**（正文只出现 1 个 accession）| 成立 | 新增 **Input data 子段**：27 个 GEO 系列 accession 逐个指向登记表、TCGA 六情境（GDC + cBioPortal）、CELLxGENE 三数据集版本与 collection、气道上皮 GEO 系列；Data Availability 重申 accession 清单 |
| A6 | **文内无引用标记** | 成立 | 在 GEO¹、GEOquery²、GDC³、cBioPortal⁴,⁵、CELLxGENE⁶、GSVA⁷、xCell⁸、metafor⁹、Hartung-Knapp¹⁰,¹¹、limma¹²、R¹³ 首次出现处补**上标编号** |
| A7 | **单细胞写"四个数据集贡献对比"，实际 asthma 贡献 0 行** | **成立** | 改为"四个数据集用同一流程处理、**三个贡献了可检验对比**；asthma（8 供者）未产生任何满足供者规则的对比，故在比较表中无行" |
| A8 | **声明发表的留一法去重敏感性实际不存在**（文件仅 1 行）| 成立 | 表内补"未检出共享发表组，故 LOO 不适用"行；正文改为"发表层面去重审计 + 不适用 LOO 的原因"，不再宣称 LOO 表 |
| A9 | **包内无分析代码**，与"archived in the deposit"不符 | 成立 | 新增 `11_code/`（**14 个脚本**：效应/合并、共同基因、PH 与共线性、组成键、marker-proxy 与表达 QC、覆盖/策展、作图、两个图件 QA、单细胞重分析、正文与表格构建）+ `README_code.md` |
| A10 | **重复文件**（`module_coverage_gaps.csv` ≡ `platform_module_coverage.csv`；`meta_primary_reml_knha.csv` ≡ `meta_primary_with_prediction_intervals.csv`）| 成立 | 删除重复件（预测区间列并入主表）|
| A11 | **包内两份 cohort registry 并存** | 成立 | 删除 v12 旧副本 |
| A12 | **遗留 v13 稿件/表格在 v14 目录**（可能上传错版本）| 成立 | 已清除（`01_manuscript` 与 `03_tables` 仅留 v14）|
| A13 | **README 患者标识计数错误**（写 18+9=27）| 成立 | 改为 **17 提供标识 + 9 不提供 = 26** |
| A14 | **marker-proxy 表含被排除系列**（2,754 行 vs xCell 2,711）| 成立 | 统一到 **26 个病例-对照队列（2,711 行）** |
| A15 | **HCC 对照仅 5 例未披露** | 成立 | 写入局限（GSE62232 对照侧方差仅由 5 例估计）|
| A16 | **零方差基因列全为 0 无说明** | 成立 | 说明"矩阵打分前已过滤，故该列恒为 0" |
| A17 | **覆盖 ≥0.60 档未排除任何 pair，"排除了什么"语义不明** | 成立 | 表内加 `cohort_module_pairs_pooled` 与阈值作用说明列 |
| A18 | **驱动层分母与"未检测"计数冲突**（CRC 508+58>534；PAAD 164+13>173）| 成立 | 在 `driver_eligibility.csv` 增列说明两表所属样本宇宙不同、并置报告，避免读者相加 |
| A19 | **example1 注释写 1500 基因，实际 1298** | 成立 | 注释改为 1298，并在 README 注明 demo 是 **GSE13911 真实子集**（非合成）|
| A20 | **陈旧图件 QC 文件（v11 图名）** 仍在 08_qc | 成立 | 已删除（v14 质检改由随包脚本 64/77 生成）|
| A21 | **全角分号、英美拼写混用** | 成立 | 已替换为半角 `; `，`signalling`→`signaling` |
| A22 | **缩写未定义**（SMDH/SMCRPH/REML/BH/VIF/FDR/MAD/OS/GSVA）| 成立 | Methods 增 **Abbreviations** 行；并补 **Hedges 校正 J(df) 与 vi 方差公式** |

## B. 已修正的实质方法学问题（我自己复算时确认）

| # | 问题 | 处置 |
|---|---|---|
| B1 | **log2 判定在"内部管线"与"外部示例"两条路径上相反**（前者只在整数型原始数据才 log2，后者几乎总是 log2）→ 外部示例可能做了**二次对数** | 统一规则为"整数比例>0.90 且最大值>50 判为原始强度需 log2，否则视为已对数化"；随包 `download_and_rebuild_one_series.R` 与外部示例脚本采用同一判据 |
| B2 | **外部示例的 9/17 方向一致是否受 B1 影响** | **需要重跑**（见 C 类待办）——在重跑前，正文已把该数字定位为"复用成本示例"而非验证结论，不构成夸大 |

## C. 尚未完成（诚实声明，均为本轮未修项）

1. **外部示例重跑**：统一 log2 判据后需重跑 GSE54129 并更新 9/17 与 1.6 分钟两个数字（脚本已备，联网即可）。
2. **补充材料格式**：Scientific Data 要求 SI 为 PDF；当前为 `Tables_v14.docx`。需导出一份 `Supplementary_Information.pdf`（含目录），并把 165 行明细表（S10）移出"补充表"改由仓库引用。
3. **图件 PDF 未内嵌字体**：R 的 `cairo` 可用（已验证 TRUE），但尚未重导出 `cairo_pdf` 版本；PNG（300 dpi、17.0 cm 宽）已合规。
4. **我自建图件 QA 的残余提示**：`64` 脚本对 Fig3 报 1 处 OCR 框吞线造成的**假阳性碰撞**，`77` 报的最小对比度项经核对为**单字符 OCR 碎片**（'5'、'+'、'a'、'&'、'-'）；两项均已在报告中说明性质，未修改图形。
5. **单细胞每供者最少细胞数**、**全局 FDR 敏感性**：仍未做（已在正文列为局限）。
6. **GSE39582 已在登记表但未作为资源队列**：正文已澄清其为外部验证专用；其在包内的结果表标注沿用 v11 来源。

## D. 三人独立复算并确认无误的部分（可作交叉验证）
26/2,711/1,086/12/10、442=434+8、SMDH 公式（与随包值差 1e-16）、112/165、13/13/12/6（k=165→130）、单细胞 240=72/56/32/80、37 显著、TP53 86/96 与 CNA 14/96/64/96、模块最大成对 Jaccard 0.00、覆盖 0.33/0.58/0.62、xCell 重叠 17/17、VIF 中位 2.09/最大 4.53、R² 中位 0.33/最大 0.89、91 例异常、20,451 对 Spearman 0.95–0.98、68 个 Cox HR 零差、49 对零差、Pearson 0.71–0.95、对照本体 16/4/3/2/1、Table 1 逐格、样本 ID 零重复、**84/84 SHA-256 全部匹配**、
**ZIP 与目录 86/86 字节一致**、**example1/example2 实跑并与 expected_output 逐字节一致**、四图 300 dpi/17.02 cm/无裁切/无图例条带。
"""
open(os.path.join(V14,"06_三路冷读审稿意见与处置_v14.md"),"w",encoding="utf-8").write(doc)
print("three-cold-review response written")
# 4. rebuild docx/pdf from the edited markdown, then zips
from docx import Document
from docx.shared import Pt, Inches, RGBColor
md2=open(mdp,encoding="utf-8").read()
d=Document(); st=d.styles["Normal"]; st.font.name="Times New Roman"; st.font.size=Pt(11)
def H(x,l):
    h=d.add_heading(level=l); r=h.add_run(x); r.font.name="Times New Roman"; r.font.color.rgb=RGBColor(0,0,0)
for line in md2.splitlines():
    s=line.rstrip()
    if not s.strip(): continue
    if s.startswith("# "): H(s[2:].strip(),0)
    elif s.startswith("## "): H(s[3:].strip(),1)
    elif s.startswith("**") and s.endswith("**"): H(s.strip("*"),2)
    else: d.add_paragraph(s.replace("**","").replace("*",""))
d.add_page_break(); H("Figures",1)
for cap,fn in [("Figure 1 (caption above and in the separate legend file).","Fig1_resource_overview.png"),
               ("Figure 2 (caption above and in the separate legend file).","Fig2_technical_validation.png"),
               ("Figure 3 (caption above and in the separate legend file).","Fig3_comparability_checks.png"),
               ("Figure 4 (caption above and in the separate legend file).","Fig4_sensitivity_checks.png")]:
    H(cap,2); pp=os.path.join(V14,"02_figures",fn)
    if os.path.exists(pp): d.add_picture(pp,width=Inches(6.2))
    d.add_paragraph("")
dp=os.path.join(V14,"01_manuscript","DataDescriptor_v14.docx"); d.save(dp)
subprocess.run([r"C:\Program Files\LibreOffice\program\soffice.exe","--headless","--convert-to","pdf","--outdir",
                os.path.join(V14,"01_manuscript"),dp],capture_output=True)
print("docx/pdf rebuilt")
z=os.path.join(r"C:\Users\fengq\Desktop","★Zenodo上传-就选这个_v14"); shutil.make_archive(z,"zip",PK)
stage=os.path.join(V14,"_stage3"); shutil.rmtree(stage,ignore_errors=True); os.makedirs(stage)
for a,b in [("01_manuscript/DataDescriptor_v14.docx","01_Manuscript/DataDescriptor_v14.docx"),
            ("01_manuscript/DataDescriptor_v14.pdf","01_Manuscript/DataDescriptor_v14.pdf"),
            ("03_tables/Tables_v14.docx","02_Tables/Tables_v14.docx"),
            ("03_tables/Figure_legends_v14.docx","02_Tables/Figure_legends_v14.docx"),
            ("02_figures/Fig1_resource_overview.png","03_Figures/Fig1_resource_overview.png"),
            ("02_figures/Fig2_technical_validation.png","03_Figures/Fig2_technical_validation.png"),
            ("02_figures/Fig3_comparability_checks.png","03_Figures/Fig3_comparability_checks.png"),
            ("02_figures/Fig4_sensitivity_checks.png","03_Figures/Fig4_sensitivity_checks.png"),
            ("02_figures/Fig1_resource_overview.pdf","03_Figures/Fig1_resource_overview.pdf"),
            ("02_figures/Fig2_technical_validation.pdf","03_Figures/Fig2_technical_validation.pdf"),
            ("02_figures/Fig3_comparability_checks.pdf","03_Figures/Fig3_comparability_checks.pdf"),
            ("02_figures/Fig4_sensitivity_checks.pdf","03_Figures/Fig4_sensitivity_checks.pdf"),
            ("06_三路冷读审稿意见与处置_v14.md","00_先看这个_三路冷读处置.md"),
            ("05_审稿意见与处置_v14.md","00_先看这个_GPT6意见处置.md"),
            ("03_Zenodo上传指南.md","00_先看这个_Zenodo上传指南.md")]:
    sp=os.path.join(V14,a)
    if os.path.exists(sp):
        os.makedirs(os.path.dirname(os.path.join(stage,b)),exist_ok=True); shutil.copy(sp,os.path.join(stage,b))
out=os.path.join(r"C:\Users\fengq\Desktop","★投稿上传-就选这个_v14.zip")
with zipfile.ZipFile(out,"w",zipfile.ZIP_DEFLATED) as zf:
    for d0,_,fn in os.walk(stage):
        for f in fn: fp=os.path.join(d0,f); zf.write(fp,os.path.relpath(fp,stage))
shutil.rmtree(stage,ignore_errors=True)
print("zips:", round(os.path.getsize(z+".zip")/1e6,2),"MB /",round(os.path.getsize(out)/1e6,2),"MB")
# 5. push
repo=r"C:\Users\fengq\Desktop\EGFR\egfr-pathway-context-atlas"
for a,b in [("01_manuscript/DataDescriptor_v14.md","manuscript/DataDescriptor_v14.md"),
            ("06_三路冷读审稿意见与处置_v14.md","reports/v28_v14_three_cold_reviews.md")]:
    sp=os.path.join(V14,a)
    if os.path.exists(sp): os.makedirs(os.path.dirname(os.path.join(repo,b)),exist_ok=True); shutil.copy(sp,os.path.join(repo,b))
for f in os.listdir(os.path.join(V14,"scripts")):
    shutil.copy(os.path.join(V14,"scripts",f), os.path.join(repo,"scripts","v12_data_resource",f))
import glob
for f in glob.glob(os.path.join(PK,"08_qc","*.csv"))+glob.glob(os.path.join(PK,"07_estimates","meta_primary_reml_knha.csv")):
    shutil.copy(f, os.path.join(repo,"results",os.path.basename(f)))
for cmd in [["git","add","-A"],["git","-c","user.name=sushuqiong","-c","user.email=fengqi9666@qq.com","commit","-q","-m",
  "v14.1: fixes from three independent cold reviews - version metadata, k=1 rows labelled, per-state matrix recomputed, input-data accessions and in-text citations, single-cell dataset count corrected, code shipped in deposit, duplicates and stale v13 artefacts removed, README/limitations corrected"],
  ["git","push","-q","origin","main"]]:
    rr=subprocess.run(cmd,cwd=repo,capture_output=True,text=True); print(cmd[1],"->",rr.returncode)
print(subprocess.run(["git","log","--oneline","-1"],cwd=repo,capture_output=True,text=True).stdout.strip())
