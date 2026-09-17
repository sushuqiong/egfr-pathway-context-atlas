#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build the v12 manuscript docx (Data Descriptor) and the Zenodo upload ZIP."""
import os, shutil, re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
V12=r"C:\Users\fengq\Desktop\EGFR\EGFR的v12"
md=open(os.path.join(V12,"01_manuscript","DataDescriptor_v12.md"),encoding="utf-8").read()
doc=Document(); st=doc.styles["Normal"]; st.font.name="Times New Roman"; st.font.size=Pt(11)
def H(x,l):
    h=doc.add_heading(level=l); r=h.add_run(x); r.font.name="Times New Roman"; r.font.color.rgb=RGBColor(0,0,0)
for line in md.splitlines():
    s=line.rstrip()
    if not s.strip(): continue
    if s.startswith("# "): H(s[2:].strip(),0)
    elif s.startswith("## "): H(s[3:].strip(),1)
    elif s.startswith("**") and s.endswith("**"): H(s.strip("*"),2)
    else: doc.add_paragraph(s.replace("**","").replace("*",""))
doc.add_page_break()
H("Figures",1)
for cap,fn in [("Figure 1. Resource overview: module gene coverage by disease context (A); assay records, unique patients and paired patients per context (B); single-cell datasets and the arms available in each (C).","Fig1_resource_overview.png"),
               ("Figure 2. Technical validation: documented exclusions, flags and non-testable comparisons (A); single-cell comparison design actually used (B); the complete estimates layer, including non-significant per-cohort effects (C).","Fig2_technical_validation.png")]:
    H(cap,2)
    p=os.path.join(V12,"02_figures",fn)
    if os.path.exists(p): doc.add_picture(p,width=Inches(6.2))
    else: doc.add_paragraph("(missing: "+fn+")")
    doc.add_paragraph("")
dp=os.path.join(V12,"01_manuscript","DataDescriptor_v12.docx"); doc.save(dp)
print("docx:",dp)
# Zenodo ZIP
zp=os.path.join(r"C:\Users\fengq\Desktop","★Zenodo上传-就选这个_v12")
shutil.make_archive(zp,"zip",os.path.join(V12,"dataset_package"))
z=zp+".zip"; print("zip:",z, round(os.path.getsize(z)/1e6,2),"MB")
# submission upload ZIP (manuscript + figures + package pointer + guides)
stage=os.path.join(V12,"_upload_stage"); shutil.rmtree(stage,ignore_errors=True); os.makedirs(stage)
for f,dst in [("01_manuscript/DataDescriptor_v12.docx","01_Manuscript/DataDescriptor_v12.docx"),
              ("01_manuscript/DataDescriptor_v12.md","01_Manuscript/DataDescriptor_v12.md"),
              ("02_figures/Fig1_resource_overview.png","02_Figures/Fig1_resource_overview.png"),
              ("02_figures/Fig2_technical_validation.png","02_Figures/Fig2_technical_validation.png"),
              ("02_figures/Fig1_resource_overview.pdf","02_Figures/Fig1_resource_overview.pdf"),
              ("02_figures/Fig2_technical_validation.pdf","02_Figures/Fig2_technical_validation.pdf"),
              ("03_tables/Tables_v12.docx","03_Tables/Tables_v12.docx"),
              ("03_tables/Figure_legends_v12.docx","03_Tables/Figure_legends_v12.docx"),
              ("03_Zenodo上传指南.md","00_先看这个_Zenodo上传指南.md"),
              ("00_状态与待决问题.md","00_先看这个_状态与待决问题.md")]:
    s=os.path.join(V12,f)
    if os.path.exists(s):
        os.makedirs(os.path.dirname(os.path.join(stage,dst)),exist_ok=True); shutil.copy(s,os.path.join(stage,dst))
out2=os.path.join(r"C:\Users\fengq\Desktop","★投稿上传-就选这个_v12.zip")
with __import__("zipfile").ZipFile(out2,"w",__import__("zipfile").ZIP_DEFLATED) as zf:
    for d0,_,fn in os.walk(stage):
        for f in fn: fp=os.path.join(d0,f); zf.write(fp,os.path.relpath(fp,stage))
print("submission zip:",out2, round(os.path.getsize(out2)/1e6,2),"MB")
shutil.rmtree(stage,ignore_errors=True)
