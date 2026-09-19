#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""v14.3 final pass: rebuild the manuscript documents from the refreshed markdown, verify, rezip, push."""
import hashlib, os, re, shutil, subprocess, zipfile
V14=r"C:\Users\fengq\Desktop\EGFR\EGFR的v14"; PK=os.path.join(V14,"dataset_package")
TBL=os.path.join(V14,"03_tables"); FIG=os.path.join(V14,"02_figures")
from docx import Document
from docx.shared import Pt, Inches, RGBColor
mp=os.path.join(V14,"01_manuscript","DataDescriptor_v14.md"); md=open(mp,encoding="utf-8").read()
CITE=re.compile(r"(GEOquery|Genomic Data Commons|cBioPortal|CELLxGENE|GEO|GSVA|xCell|metafor|Hartung-Knapp|limma|R)(\d+(?:,\d+)*)")
def add_para(doc,text):
    p=doc.add_paragraph(); pos=0
    for m in CITE.finditer(text):
        if m.start()>pos: p.add_run(text[pos:m.start()])
        p.add_run(m.group(1)); s=p.add_run(m.group(2)); s.font.superscript=True; pos=m.end()
    if pos<len(text): p.add_run(text[pos:])
    for r in p.runs: r.font.name="Times New Roman"
doc=Document(); st=doc.styles["Normal"]; st.font.name="Times New Roman"; st.font.size=Pt(11)
figs=[("Figure 1 (legend in the separate figure-legend file).","Fig1_resource_overview.png"),
      ("Figure 2 (legend in the separate figure-legend file).","Fig2_technical_validation.png"),
      ("Figure 3 (legend in the separate figure-legend file).","Fig3_comparability_checks.png"),
      ("Figure 4 (legend in the separate figure-legend file).","Fig4_sensitivity_checks.png")]
for line in md.splitlines():
    s=line.rstrip()
    if not s.strip(): continue
    if s.startswith("# "):
        h=doc.add_heading(level=0); r=h.add_run(s[2:].strip()); r.font.name="Times New Roman"; r.font.color.rgb=RGBColor(0,0,0)
    elif s.startswith("## "):
        h=doc.add_heading(level=1); r=h.add_run(s[3:].strip()); r.font.name="Times New Roman"; r.font.color.rgb=RGBColor(0,0,0)
    elif s.startswith("**") and s.endswith("**") and not s.startswith("**Figure"):
        h=doc.add_heading(level=2); r=h.add_run(s.strip("*")); r.font.name="Times New Roman"; r.font.color.rgb=RGBColor(0,0,0)
    else: add_para(doc, s.replace("**","").replace("*",""))
doc.add_page_break(); h=doc.add_heading(level=1); r=h.add_run("Figures"); r.font.name="Times New Roman"; r.font.color.rgb=RGBColor(0,0,0)
for cap,fn in figs:
    h=doc.add_heading(level=2); r=h.add_run(cap); r.font.name="Times New Roman"; r.font.color.rgb=RGBColor(0,0,0)
    doc.add_picture(os.path.join(FIG,fn), width=Inches(6.2)); doc.add_paragraph("")
dp=os.path.join(V14,"01_manuscript","DataDescriptor_v14.docx"); doc.save(dp)
subprocess.run([r"C:\Program Files\LibreOffice\program\soffice.exe","--headless","--convert-to","pdf","--outdir",
                os.path.join(V14,"01_manuscript"),dp],capture_output=True)
d=Document(dp); dt="\n".join(p.text for p in d.paragraphs)
disk={fn:hashlib.sha256(open(os.path.join(FIG,fn),"rb").read()).hexdigest() for _,fn in figs}
em=[hashlib.sha256(r.target_part.blob).hexdigest() for r in d.part.rels.values() if "image" in r.reltype]
print("docx: files-count phrase present:", "107 files" in dt, "| S17 present:", "Table S17" in dt,
      "| figures matching disk:", f"{sum(1 for h in em if h in disk.values())}/{len(em)}",
      "| superscript runs:", sum(1 for p in d.paragraphs for r in p.runs if r.font.superscript))
# rezip the submission package with the fresh documents
z=os.path.join(r"C:\Users\fengq\Desktop","★Zenodo上传-就选这个_v14"); shutil.make_archive(z,"zip",PK)
stage=os.path.join(V14,"_stage7"); shutil.rmtree(stage,ignore_errors=True)
pairs=[("01_manuscript/DataDescriptor_v14.docx","01_Manuscript/DataDescriptor_v14.docx"),
       ("01_manuscript/DataDescriptor_v14.pdf","01_Manuscript/DataDescriptor_v14.pdf"),
       ("03_tables/Supplementary_Information.pdf","02_Supplementary/Supplementary_Information.pdf"),
       ("03_tables/Tables_v14.docx","02_Supplementary/Tables_v14.docx"),
       ("03_tables/Figure_legends_v14.docx","02_Supplementary/Figure_legends_v14.docx")]
names={1:"resource_overview",2:"technical_validation",3:"comparability_checks",4:"sensitivity_checks"}
for i in (1,2,3,4):
    for ext in ("png","pdf"): pairs.append((f"02_figures/Fig{i}_{names[i]}.{ext}", f"03_Figures/Fig{i}_{names[i]}.{ext}"))
for a,b in pairs:
    sp=os.path.join(V14,a)
    if os.path.exists(sp):
        os.makedirs(os.path.dirname(os.path.join(stage,b)),exist_ok=True); shutil.copy(sp,os.path.join(stage,b))
open(os.path.join(stage,"00_READ_ME_FIRST.txt"),"w",encoding="utf-8").write(
"Submission package - cross-disease transcriptomic resource (v14)\n\n"
"01_Manuscript : DataDescriptor_v14.docx (editable) and .pdf\n"
"02_Supplementary : Supplementary_Information.pdf (Table 1 and Supplementary Tables S1-S17) and the same tables as .docx\n"
"03_Figures : Figures 1-4 as .png (300 dpi) and .pdf\n\n"
"The data deposit (including analysis code) is provided separately as the Zenodo archive; the DOI is inserted in the\n"
"manuscript once the deposit is published. Internal working notes are intentionally not part of this package.\n")
out=os.path.join(r"C:\Users\fengq\Desktop","★投稿上传-就选这个_v14.zip")
with zipfile.ZipFile(out,"w",zipfile.ZIP_DEFLATED) as zf:
    for d0,_,fn in os.walk(stage):
        for f in fn: fp=os.path.join(d0,f); zf.write(fp,os.path.relpath(fp,stage))
shutil.rmtree(stage,ignore_errors=True)
print("zips:", round(os.path.getsize(z+".zip")/1e6,2),"MB /",round(os.path.getsize(out)/1e6,2),"MB")
repo=r"C:\Users\fengq\Desktop\EGFR\egfr-pathway-context-atlas"
shutil.copy(mp, os.path.join(repo,"manuscript","DataDescriptor_v14.md"))
for f in os.listdir(os.path.join(V14,"scripts")):
    src=os.path.join(V14,"scripts",f)
    if os.path.isfile(src): shutil.copy(src, os.path.join(repo,"scripts","v12_data_resource",f))
for f in ["single_cell_sensitivity_min10cells_per_donor.csv","single_cell_multiplicity_strategy_comparison.csv"]:
    p=os.path.join(PK,"06_single_cell",f)
    if os.path.exists(p): shutil.copy(p, os.path.join(repo,"results",f))
for f in ["figure_qa_notes.md","checksums_sha256.txt"]:
    p=os.path.join(PK,"08_qc",f)
    if os.path.exists(p): shutil.copy(p, os.path.join(repo,"results",f))
for cmd in [["git","add","-A"],["git","-c","user.name=sushuqiong","-c","user.email=fengqi9666@qq.com","commit","-q","-m",
            "v14.3: manuscript documents rebuilt from the final markdown (counts regenerated from the manifest, byte-identical embedded figures, superscript citations), submission package free of internal notes"],
            ["git","push","-q","origin","main"]]:
    rr=subprocess.run(cmd,cwd=repo,capture_output=True,text=True); print(cmd[1],"->",rr.returncode)
print(subprocess.run(["git","log","--oneline","-1"],cwd=repo,capture_output=True,text=True).stdout.strip())
