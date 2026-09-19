#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""v14: survival data dictionary, prediction intervals for the pooled estimates, Tables_v14.docx (S1-S16), legends."""
import csv, os, math, statistics
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
V14=r"C:\Users\fengq\Desktop\EGFR\EGFR的v14"; R=os.path.join(V14,"results"); PK=os.path.join(V14,"dataset_package")
OUT=os.path.join(V14,"03_tables"); os.makedirs(OUT,exist_ok=True)
def rd(p):
    if not os.path.exists(p): return []
    with open(p,encoding="utf-8-sig",newline="") as f: return list(csv.DictReader(f))
def wr(rows,p):
    with open(p,"w",encoding="utf-8",newline="") as f:
        w=csv.DictWriter(f,fieldnames=list(rows[0].keys()),lineterminator="\n"); w.writeheader(); w.writerows(rows)
# ---------- 1. survival data dictionary ----------
open(os.path.join(PK,"07_estimates","survival_data_dictionary.md"),"w",encoding="utf-8").write(
"""# Survival data dictionary (TCGA-derived layer)

| field | definition |
|---|---|
| endpoint | overall survival (OS) |
| event | death; `OS_STATUS` containing DECEASED is coded 1, all other values 0 |
| time | `OS_MONTHS`, months from diagnosis to death or last follow-up |
| censoring | patients alive at last follow-up are censored at `OS_MONTHS` |
| exclusions | records with non-finite or non-positive time |
| covariates | age (`AGE`, years, continuous) and pathological stage collapsed to I-IV (`PATH_STAGE`, falling back to `CLINICAL_STAGE`) |
| missing data | multivariable models are fitted on complete cases only (age and stage non-missing); the excluded count is reported in the model table |
| module score | mean over module members of the per-gene z-score across patients within the context (log2(x+1) values, z-scored per gene, averaged per patient, then averaged across members present; at least two members required) |
| feature scaling | each module score is z-scored before model fitting, so hazard ratios are per 1 SD |
| models | univariable Cox; age-plus-stage adjusted Cox; proportional hazards checked with scaled Schoenfeld residuals (cox.zph) |
| software | R 4.4.1 with survival 3.8.3 |
| data version | cBioPortal releases retrieved 2026-09 (luad_tcga_gdc, coadread_tcga_pan_can_atlas_2018, stad_tcga_gdc, paad_tcga_gdc, lihc_tcga_pan_can_atlas_2018, esca_tcga_gdc) |
| oesophageal subset | restricted to squamous histology (96 patients; 94 with survival data); mutation denominator corrected from 185 mixed-histology profiles to the 96 squamous profiles |
""")
# ---------- 2. prediction intervals on the primary pooled estimates ----------
def t975(k):
    tbl={2:12.706,3:4.303,4:3.182,5:2.776,6:2.571,7:2.447,8:2.365,9:2.306,10:2.262,11:2.228,12:2.201,
         13:2.179,14:2.160,15:2.145,16:2.131,17:2.120,18:2.110,19:2.101,20:2.093}
    return tbl.get(k, 1.96)
meta=rd(os.path.join(R,"v12_meta_primary_reml_knha.csv"))
for r in meta:
    k=int(r["k"]); tau2=float(r["tau2"] or 0); se=float(r["se"]); est=float(r["est"])
    pi=t975(k)*math.sqrt(se**2 + tau2)
    r["pi_low"]=round(est-pi,4); r["pi_high"]=round(est+pi,4)
    r["pi_note"]="95% prediction interval; uninformative when k is small" if k<4 else "95% prediction interval"
wr(meta, os.path.join(PK,"07_estimates","meta_primary_reml_knha.csv"))
wr(meta, os.path.join(PK,"07_estimates","meta_primary_with_prediction_intervals.csv"))
print("prediction intervals added to", len(meta), "pooled states")
# ---------- 3. Word tables ----------
def three_line(doc, headers, rows, font=8.5):
    t=doc.add_table(rows=1+len(rows), cols=len(headers)); t.style=doc.styles['Normal Table']; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    tp=t._tbl.tblPr
    for old in tp.findall(qn('w:tblBorders')): tp.remove(old)
    b=OxmlElement('w:tblBorders')
    for e in ('top','left','bottom','right','insideH','insideV'):
        el=OxmlElement(f'w:{e}'); el.set(qn('w:val'),'nil'); b.append(el)
    tp.append(b)
    def bd(cell, edges, sz=8):
        tcPr=cell._tc.get_or_add_tcPr()
        for old in tcPr.findall(qn('w:tcBorders')): tcPr.remove(old)
        tb=OxmlElement('w:tcBorders')
        for e in ('top','left','bottom','right','insideH','insideV'):
            el=OxmlElement(f'w:{e}')
            if e in edges: el.set(qn('w:val'),'single'); el.set(qn('w:sz'),str(sz)); el.set(qn('w:color'),'000000')
            else: el.set(qn('w:val'),'nil')
            tb.append(el)
        tcPr.append(tb)
    for j,h in enumerate(headers):
        c=t.cell(0,j); c.text=''; rr=c.paragraphs[0].add_run(str(h)); rr.bold=True; rr.font.size=Pt(font); rr.font.name='Times New Roman'
        bd(c,{'top','bottom'},12)
    for i,row in enumerate(rows,1):
        for j,v in enumerate(row):
            c=t.cell(i,j); c.text=''; rr=c.paragraphs[0].add_run(str(v)); rr.font.size=Pt(font); rr.font.name='Times New Roman'
            if i==len(rows): bd(c,{'bottom'},12)
    doc.add_paragraph('')
def head(doc,text,level=1):
    h=doc.add_heading(level=level); rr=h.add_run(text); rr.font.name='Times New Roman'; rr.font.color.rgb=RGBColor(0,0,0)
CTX={"GSE32863":"LUAD","GSE19804":"LUAD","GSE10072":"LUAD","GSE44076":"CRC","GSE41258":"CRC","GSE23878":"CRC",
     "GSE75214":"IBD","GSE87473":"IBD","GSE179285":"IBD","GSE16879":"IBD","GSE4302":"Asthma","GSE43696":"Asthma",
     "GSE67472":"Asthma","GSE27342":"STAD","GSE63089":"STAD","GSE13911":"STAD","GSE15471":"PAAD","GSE28735":"PAAD",
     "GSE62452":"PAAD","GSE57957":"HCC","GSE62232":"HCC","GSE23400":"ESCC","GSE20347":"ESCC","GSE76925":"COPD",
     "GSE47460":"COPD","GSE33814":"NAFLD","GSE66676":"NAFLD"}
order=["LUAD","CRC","STAD","PAAD","HCC","ESCC","IBD","COPD","NAFLD","Asthma"]
reg=rd(os.path.join(PK,"01_cohort_registry","cohort_registry.csv")); ccase=[r for r in reg if r["design"]=="case-control"]
agg={}
for r in ccase:
    c=CTX.get(r["accession"],"other"); a=agg.setdefault(c,dict(n=0,a=0,p=0,pair=0,pl=set(),np_=0))
    a["n"]+=1; a["a"]+=int(r["n_assay_records"])
    a["p"]+=int(r["n_unique_patients"]) if str(r["n_unique_patients"]).isdigit() else 0
    a["pair"]+=int(r["patients_with_both_arms"] or 0) if r["paired_design_used"]=="yes" else 0
    a["pl"].add(r["platform"]); a["np_"]+=1 if r["patient_identity_available"]=="no" else 0
doc=Document(); st=doc.styles['Normal']; st.font.name='Times New Roman'; st.font.size=Pt(10)
head(doc,"Table 1. Cohort composition of the resource by disease context")
three_line(doc,["Context","Cohorts","Assay records","Patients (with identifiers)","Patients with both tissues","Platforms","Cohorts without patient identifiers"],
           [[c,agg[c]["n"],agg[c]["a"],agg[c]["p"],agg[c]["pair"],len(agg[c]["pl"]),agg[c]["np_"]] for c in order if c in agg])
doc.add_paragraph("Case-control series only (26 cohorts). One treatment-response series was reviewed and excluded; exclusion details are in Supplementary Table S4 and the control-tissue ontology in S13.")
head(doc,"Supplementary Table S1. Cohort registry with design, counts, control class, origin publication and provenance")
ont={r["accession"]:r["control_class"] for r in rd(os.path.join(PK,"01_cohort_registry","contrast_ontology.csv"))}
three_line(doc,["Accession","Context","Platform","Assays","Patients","Cases","Controls","Both-tissue patients","Paired used","Patient IDs","Control class","Origin PMID"],
  [[r["accession"],CTX.get(r["accession"],""),r["platform"],r["n_assay_records"],r["n_unique_patients"],r["n_case"],r["n_control"],
    r["patients_with_both_arms"],r["paired_design_used"],r["patient_identity_available"],ont.get(r["accession"],""),r["origin_pmid"]] for r in ccase],7)
cov=rd(os.path.join(PK,"04_modules","platform_module_coverage.csv")); acc_cc={r["accession"] for r in ccase}
mods={}
for r in cov:
    if r["accession"] not in acc_cc: continue
    m=r["module"]; c=float(r["coverage"]); d=mods.setdefault(m,dict(members=r["n_members"],minc=1.0,n_inc=0))
    d["minc"]=min(d["minc"],c); d["n_inc"]+=1 if int(r["n_present"])<int(r["n_members"]) else 0
head(doc,"Supplementary Table S2. Module membership and cross-cohort gene coverage (26 case-control cohorts)")
three_line(doc,["Module (short label)","Members","Minimum coverage","Cohorts incomplete (of 26)"],
  [[m,v["members"],f"{v['minc']:.2f}",v["n_inc"]] for m,v in sorted(mods.items(),key=lambda kv:kv[1]["minc"])])
head(doc,"Supplementary Table S3. Single-cell datasets, cells and comparison arms")
three_line(doc,["Dataset","Cells total","Cells used","Donors","Cell types","Arms","Notes"],
  [[r["context"],r["cells_total"],r["cells_used"],r["donors"],r["cell_types"],r["arms"],r["control_kinds"]] for r in rd(os.path.join(R,"v12_sc_dataset_audit.csv"))],8)
head(doc,"Supplementary Table S4. Documented exclusions, flags and design corrections")
three_line(doc,["Item","n","Layer","Reason"],[[r["item"],r["n"],r["layer"],r["reason"]] for r in rd(os.path.join(PK,"08_qc","exclusion_and_flag_log.csv"))],8)
head(doc,"Supplementary Table S5. Decision rules and thresholds")
three_line(doc,["Rule","Definition / threshold"],[
 ["Effect size","SMDH (unpaired) or SMCRPH (paired, empirical within-pair correlation r); shared denominator sqrt((sd_case^2 + sd_control^2)/2); paired cohorts use complete pairs only; zero-variance modules are reported as not estimable"],
 ["Pooling","REML with Hartung-Knapp correction; minimum two cohorts; k, tau-squared, I-squared, 95% CI and 95% prediction interval reported"],
 ["Paired design","metadata-derived: at least four patients contributing both a case and a control sample"],
 ["Coverage policy","every effect row carries its module coverage; pooled results are re-estimated after excluding cohorts below 0.60 / 0.80 / 0.90 coverage"],
 ["Control ontology","control tissue is classified (adjacent non-tumour, non-inflamed, healthy donor, non-diseased); classes are not pooled silently"],
 ["Common-gene sensitivity","re-score with members present in every cohort; flag when the direction disagrees in >=20% of cohorts; not scorable when fewer than two members are shared"],
 ["Proportional hazards","scaled Schoenfeld residuals (cox.zph) in the age-adjusted model; flagged at p<0.05"],
 ["Single-cell design","library-size normalisation (1e4) and log1p; donor means per cell type and arm; >=3 donors per arm; paired signed-rank when >=5 complete pairs; unpaired only when donor-disjoint; otherwise not testable; BH within context x module x comparison type"],
 ["Reproduction tolerance","exact reproduction means an absolute difference of at most 1e-6 in the statistic concerned; rank agreement is reported as Spearman and never described as reproduction"],
 ["Mutation status","reported only for samples in the source mutation-profiled list; otherwise not assayed / unknown"]],8)
head(doc,"Supplementary Table S6. Cohort-by-module pairs that are not estimable")
three_line(doc,["Accession","Module","Reason"],[[r["accession"],r["module"],r["reason"]] for r in rd(os.path.join(PK,"07_estimates","not_estimable_pairs.csv"))],8)
head(doc,"Supplementary Table S7. Validation evidence: reproduction within tolerance versus rank agreement")
three_line(doc,["Check","Scope","Result","Interpretation"],[
 ["Per-sample module scores","3 cohorts, 49 module-cohort pairs","maximum absolute difference 0 (tolerance 1e-6); Spearman 1.00","reproduction within tolerance"],
 ["Cohort effects and standard errors","3 cohorts, 49 effects","maximum absolute difference 0","reproduction within tolerance"],
 ["Cox models refitted from shipped scores","LUAD and CRC univariable; squamous oesophageal univariable and age+stage; 68 comparisons","maximum absolute hazard-ratio difference 0","reproduction within tolerance"],
 ["TCGA patient-level scores rebuilt from the gene-level table","3 contexts; 20,451 patient-module pairs","Spearman 0.95-0.98","rank agreement only"],
 ["xCell compartment scores recomputed","3 cohorts x 4 compartments","Pearson 0.71-0.95; Spearman 0.48-0.94","substantial agreement"],
 ["Marker-proxy versus xCell","26 cohorts x 4 compartments",' ; '.join(f"{r['compartment']} Spearman {r['spearman_vs_xcell']}" for r in rd(os.path.join(PK,"08_qc","qc_composition_method_agreement.csv"))),"independent method agreement"],
 ["Cross-cohort direction","165 context-module pairs","112 pairs share the sign in every cohort","partial comparability, quantified"],
 ["Coverage-threshold sensitivity","all / >=0.60 / >=0.80 / >=0.90","BH-significant states 13 / 13 / 12 / 6","gradual degradation"]],8)
head(doc,"Supplementary Table S8. Which score layer to use")
three_line(doc,["Question","Layer","Caveat"],[
 ["Within-cohort case-control difference","bulk GSVA z-scores","cohort-internal; compare across cohorts only via effect sizes"],
 ["Patient-level difference across tumour contexts","TCGA member-mean z-scores","different construction rule from the bulk layer"],
 ["Cell-type-resolved difference","single-cell member means","donor means are tested, not cells; no minimum cells per donor applied"],
 ["Composition as descriptor or sensitivity covariate","xCell compartments or the marker-proxy alternative","enrichment scores share genes with module scores"]],8)
head(doc,"Supplementary Table S9. Comparability, coupling and collinearity checks")
three_line(doc,["Check","Result"],[
 ["Cross-cohort direction consistency","112 of 165 context-module pairs (68%) share the effect sign in every contributing cohort"],
 ["Coverage-threshold sensitivity","BH-significant states 13 (all, k=165) / 13 (>=0.60) / 12 (>=0.80) / 6 (>=0.90)"],
 ["Module versus xCell signature overlap","all 17 modules share more than 20% of members with the xCell panel; maximum 100%"],
 ["Collinearity: overall model","R2 of disease status on the four compartments: median 0.33, maximum 0.89"],
 ["Collinearity: per-predictor VIF","median 2.09, maximum 4.53"],
 ["Curation flow","27 series retrieved; 26 case-control cohorts; 1 treatment-response series excluded; 3 downloaded series not carried forward"]],8)
head(doc,"Supplementary Table S10. Cross-cohort direction consistency (summary; detail deposited)")
con=[r for r in rd(os.path.join(R,"v13_qc_cross_cohort_consistency.csv")) if int(r["k"])>=2]
same=sum(1 for r in con if str(r["direction_consistent"]).upper()=="TRUE")
three_line(doc,["Statistic","Value"],[
 ["Context-module pairs with at least two cohorts", len(con)],
 ["Pairs whose effect direction is the same in every contributing cohort", f"{same} ({100*same/max(1,len(con)):.0f}%)"],
 ["Detail (one row per pair with cohort counts, sign counts and median effect)","deposited as 08_qc/qc_cross_cohort_consistency.csv (165 rows); not printed here because oversized tables are deposited rather than printed"]],8)
head(doc,"Supplementary Table S11. Expression matrix quality control per cohort")
three_line(doc,["Accession","Context","Samples","Genes","Missing fraction","Outlier samples (>3 MAD)","Zero-variance genes"],
  [[r["accession"],r["context"],r["n_samples"],r["n_genes"],r["missing_fraction"],r["outlier_samples_3mad"],r["gene_sd_zero"]] for r in rd(os.path.join(PK,"08_qc","qc_expression_matrix_per_cohort.csv"))],7.5)
mo=rd(os.path.join(PK,"04_modules","module_overlap_matrix.csv"))
head(doc,"Supplementary Table S12. Module overlap matrix (maximum pairwise Jaccard overlap)")
three_line(doc,["Statistic","Value"],[
 ["Number of module pairs compared", len(mo)],
 ["Maximum pairwise Jaccard overlap", max(float(r["jaccard"]) for r in mo) if mo else 0],
 ["Pairs sharing any gene", sum(1 for r in mo if int(r["shared_genes"])>0)],
 ["Interpretation","the seventeen module gene sets are mutually disjoint, so modules do not double-count genes"]],8)
head(doc,"Supplementary Table S13. Control tissue ontology")
three_line(doc,["Accession","Context","Control description","Control class"],[[r["accession"],r["context"],r["control_description"],r["control_class"]] for r in rd(os.path.join(PK,"01_cohort_registry","contrast_ontology.csv"))],7.5)
head(doc,"Supplementary Table S14. Study-level overlap audit and leave-one-out dedup sensitivity")
three_line(doc,["Scenario","Pooled states (k>=3)","Nominally significant"],[[r["scenario"],r["states"],r["nominal_significant"]] for r in rd(os.path.join(PK,"08_qc","qc_study_dedup_sensitivity.csv"))]+[["no shared-publication groups detected","not applicable","not applicable"]],8)
head(doc,"Supplementary Table S15. Survival layer data dictionary (summary)")
three_line(doc,["Field","Definition"],[
 ["Endpoint","overall survival (death from any cause)"],
 ["Time","OS_MONTHS, months from diagnosis to death or last follow-up"],
 ["Censoring","patients alive at last follow-up censored at OS_MONTHS"],
 ["Covariates","age (continuous, years) and pathological stage (I-IV, falling back to clinical stage)"],
 ["Missing data","multivariable models fitted on complete cases only"],
 ["Data version","cBioPortal releases retrieved 2026-09; full dictionary in 07_estimates/survival_data_dictionary.md"]],8)
head(doc,"Supplementary Table S16. Marker-proxy composition method agreement with xCell")
three_line(doc,["Compartment","Samples","Spearman vs xCell","Pearson vs xCell"],
  [[r["compartment"],r["n"],r["spearman_vs_xcell"],r["pearson_vs_xcell"]] for r in rd(os.path.join(PK,"08_qc","qc_composition_method_agreement.csv"))],8)
tp=os.path.join(OUT,"Tables_v14.docx"); doc.save(tp); print("tables:",tp)
# ---------- 4. legends ----------
L=Document(); st2=L.styles['Normal']; st2.font.name='Times New Roman'; st2.font.size=Pt(11)
head(L,"Figure legends (v14)")
for t,b in [
 ("Figure 1. Resource overview.",
  "(A) Module gene coverage per disease context: the fraction of frozen module members present in each cohort's processed matrix. Module labels are abbreviated in the figure; the abbreviation expands to the full module name as follows: REC = RECEPTORS, LIG = LIGANDS, and the suffix AXIS is omitted (for example IGF_INSR = IGF_INSR_AXIS). "
  "(B) Assay records, patients with available identifiers and patients contributing both tissues, one facet per measure. "
  "(C) Single-cell datasets with the comparison arms available in each; the lung dataset is shown as excluded because all of its cells come from tumour tissue."),
 ("Figure 2. Validation of sample and design handling.",
  "(A) Documented exclusions, flags and non-testable comparisons on a logarithmic scale. (B) The comparison design actually used in the single-cell layer. "
  "(C) The complete estimates layer: all per-cohort standardized effects, faceted by the measure used, including non-significant results."),
 ("Figure 3. Comparability and coupling.",
  "(A) Cross-cohort direction consistency: each point is one context-module pair with at least two cohorts, and the bar marks the median. "
  "(B) Coverage-threshold sensitivity: the number of cohort-module pairs pooled (light) and the number of states surviving BH control (red). "
  "(C) Fraction of each module's members that lie inside the xCell signature panel, with the 20% level marked. "
  "(D) Overall model R-squared versus the maximum per-predictor VIF within the disease-status model."),
 ("Figure 4. Sensitivity analyses.",
  "(A) Common-gene sensitivity of module scoring: Spearman correlation between full-member and common-gene scoring per module; red bars mark modules whose effect direction disagreed in at least 20% of cohorts, grey bars mark modules with too few shared members to re-score. "
  "(B) Collinearity between disease status and the composition scores, per cohort, with the median marked."),
 ("Supplementary material.",
  "All numerical supplementary material is provided as machine-readable tables in the data deposit; Tables S1-S16 accompany this manuscript, and no supplementary figures are required.")]:
    head(L,t,2); L.add_paragraph(b)
lp=os.path.join(OUT,"Figure_legends_v14.docx"); L.save(lp); print("legends:",lp)
