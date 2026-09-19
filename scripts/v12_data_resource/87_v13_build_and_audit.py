#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""v13 deliverables: manuscript docx/pdf, submission and Zenodo ZIPs, v13 audits."""
import csv, os, re, shutil, subprocess, sys, zipfile
V13=r"C:\Users\fengq\Desktop\EGFR\EGFR的v13"; R=os.path.join(V13,"results"); PK=os.path.join(V13,"dataset_package")
def rd(p):
    with open(p,encoding="utf-8-sig",newline="") as f: return list(csv.DictReader(f))
# ---------- 1. manuscript docx ----------
from docx import Document
from docx.shared import Pt, Inches, RGBColor
md=open(os.path.join(V13,"01_manuscript","DataDescriptor_v13.md"),encoding="utf-8").read()
doc=Document(); st=doc.styles["Normal"]; st.font.name="Times New Roman"; st.font.size=Pt(11)
def H(x,l):
    h=doc.add_heading(level=l); r=h.add_run(x); r.font.name="Times New Roman"; r.font.color.rgb=RGBColor(0,0,0)
for line in md.splitlines():
    s=line.rstrip()
    if not s.strip(): continue
    if s.startswith("# "): H(s[2:].strip(),0)
    elif s.startswith("## "): H(s[3:].strip(),1)
    elif s.startswith("**") and s.endswith("**"): H(s.strip("*"),2)
    else: doc.add_paragraph(s.replace("**","").replace("*",""))
doc.add_page_break(); H("Figures",1)
caps=[("Figure 1 (caption in the text above and in the separate legend file).","Fig1_resource_overview.png"),
      ("Figure 2 (caption in the text above and in the separate legend file).","Fig2_technical_validation.png"),
      ("Figure 3 (caption in the text above and in the separate legend file).","Fig3_comparability_checks.png")]
for cap,fn in caps:
    H(cap,2); p=os.path.join(V13,"02_figures",fn)
    if os.path.exists(p): doc.add_picture(p,width=Inches(6.2))
    else: doc.add_paragraph("(missing: "+fn+")")
dp=os.path.join(V13,"01_manuscript","DataDescriptor_v13.docx"); doc.save(dp); print("docx:",dp)
subprocess.run([r"C:\Program Files\LibreOffice\program\soffice.exe","--headless","--convert-to","pdf","--outdir",
                os.path.join(V13,"01_manuscript"),dp],capture_output=True)
print("pdf:", os.path.exists(os.path.join(V13,"01_manuscript","DataDescriptor_v13.pdf")))
# ---------- 2. ZIPs ----------
z=os.path.join(r"C:\Users\fengq\Desktop","★Zenodo上传-就选这个_v13")
shutil.make_archive(z,"zip",PK); print("zenodo zip:",z+".zip",round(os.path.getsize(z+".zip")/1e6,2),"MB")
stage=os.path.join(V13,"_stage"); shutil.rmtree(stage,ignore_errors=True); os.makedirs(stage)
files=[("01_manuscript/DataDescriptor_v13.docx","01_Manuscript/DataDescriptor_v13.docx"),
       ("01_manuscript/DataDescriptor_v13.pdf","01_Manuscript/DataDescriptor_v13.pdf"),
       ("01_manuscript/DataDescriptor_v13.md","01_Manuscript/DataDescriptor_v13.md"),
       ("03_tables/Tables_v13.docx","02_Tables/Tables_v13.docx"),
       ("03_tables/Figure_legends_v13.docx","02_Tables/Figure_legends_v13.docx"),
       ("02_figures/Fig1_resource_overview.png","03_Figures/Fig1_resource_overview.png"),
       ("02_figures/Fig2_technical_validation.png","03_Figures/Fig2_technical_validation.png"),
       ("02_figures/Fig3_comparability_checks.png","03_Figures/Fig3_comparability_checks.png"),
       ("02_figures/Fig1_resource_overview.pdf","03_Figures/Fig1_resource_overview.pdf"),
       ("02_figures/Fig2_technical_validation.pdf","03_Figures/Fig2_technical_validation.pdf"),
       ("02_figures/Fig3_comparability_checks.pdf","03_Figures/Fig3_comparability_checks.pdf"),
       ("03_Zenodo上传指南.md","00_先看这个_Zenodo上传指南.md"),
       ("00_状态与可继续事项_v13.md","00_先看这个_状态与可继续事项.md")]
for s,d in files:
    sp=os.path.join(V13,s)
    if os.path.exists(sp):
        os.makedirs(os.path.dirname(os.path.join(stage,d)),exist_ok=True); shutil.copy(sp,os.path.join(stage,d))
out=os.path.join(r"C:\Users\fengq\Desktop","★投稿上传-就选这个_v13.zip")
with zipfile.ZipFile(out,"w",zipfile.ZIP_DEFLATED) as zf:
    for d0,_,fn in os.walk(stage):
        for f in fn: fp=os.path.join(d0,f); zf.write(fp,os.path.relpath(fp,stage))
print("submission zip:",out,round(os.path.getsize(out)/1e6,2),"MB")
shutil.rmtree(stage,ignore_errors=True)
# ---------- 3. audits ----------
checks=[]; issues=[]
def chk(c,g,b=None):
    (checks if c else issues).append(g if c else (b or g))
cc=rd(os.path.join(PK,"01_cohort_registry","cohort_registry.csv")); cc=[r for r in cc if r["design"]=="case-control"]
mut=rd(os.path.join(R,"v12_escc_mutation_denominators.csv")); cn=rd(os.path.join(R,"v12_escc_cna_summary.csv"))
tp53=[r for r in mut if r["gene"]=="TP53"]; egfr=[r for r in cn if r["gene"]=="EGFR"]
full=rd(os.path.join(PK,"07_estimates","per_cohort_effects.csv"))
scc=rd(os.path.join(R,"v12_sc_comparisons.csv")); ph=rd(os.path.join(R,"v12_qc_cox_ph_assumption.csv"))
title=md.splitlines()[0].lstrip("# ").strip(); ab=re.search(r"\*\*Abstract\*\*(.*?)\*\*Background", md, re.S).group(1)
abw=len(re.findall(r"\S+",ab))
chk(len(title)<=110, f"title {len(title)} chars"); chk(abw<=170, f"abstract {abw} words")
chk(len(tp53)==1 and tp53[0]["n_mutated_samples"]=="86" and "86" in md and "TP53 is reported mutated in 86" in md,
    "TP53 reported as 86/96 with an explicit gene lookup", "TP53 value missing or mislabelled in the manuscript")
chk("TP53 reported in 0" not in md, "no stale 'TP53 reported in 0' phrase", "stale TP53 phrase present")
chk(len(full)==26*17, f"estimates matrix complete: {len(full)} = 26 x 17 rows", f"estimates matrix has {len(full)} rows")
chk(sum(1 for r in full if r["measure"]=="not estimable")==8, "8 not-estimable rows retained with reasons", "not-estimable rows missing")
chk(all(r.get("not_estimable_reason") for r in full if r["measure"]=="not estimable"), "every not-estimable row carries a reason")
chk(len(scc)==240, f"single-cell comparisons: {len(scc)}")
chk(sum(1 for r in ph if r["violation" if "violation" in r else "ph_violation"].upper()=="TRUE")==0, "no PH violations")
for f in ["00_VERSION.txt","00_CHANGELOG.md","08_qc/curation_flow.csv","08_qc/qc_cross_cohort_consistency.csv",
          "08_qc/qc_coverage_threshold_sensitivity.csv","08_qc/qc_module_vs_xcell_signature_overlap.csv",
          "08_qc/qc_vif_by_predictor.csv","07_estimates/not_estimable_pairs.csv"]:
    chk(os.path.exists(os.path.join(PK,f)), f"package file {f}", f"MISSING {f}")
for f in ["Tables_v13.docx","Figure_legends_v13.docx"]:
    chk(os.path.exists(os.path.join(V13,"03_tables",f)), f"Word output {f}", f"MISSING {f}")
for tok in ["TODO","TBD","XXX","YOUR_MATRIX","待定"]:
    chk(tok.lower() not in md.lower(), f"no placeholder {tok}", f"placeholder {tok}")
figs=[f for f in os.listdir(os.path.join(V13,"02_figures")) if f.endswith(".png")]
chk(len(figs)==3, f"three main figures ({len(figs)})", f"expected 3 figures, found {len(figs)}")
for i in (1,2,3): chk(f"Figure {i}" in md, f"Figure {i} cited")
inv=rd(os.path.join(PK,"08_qc","file_inventory_and_checksums.csv"))
chk(len(inv)>=60, f"inventory covers {len(inv)} files", "inventory too small")
print("\n=== v13 AUDIT ===")
print("passed:",len(checks)); print("issues:",len(issues))
for i in issues: print("  ISSUE",i)
print("figures:",sorted(figs))
print("docx/pdf/zip built; package files:",len(inv))
