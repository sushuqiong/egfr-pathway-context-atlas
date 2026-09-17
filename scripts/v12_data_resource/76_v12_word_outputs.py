#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""v12: build (a) Tables_v12.docx with booktabs-style three-line tables and (b) Figure_legends_v12.docx."""
import csv, os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
V12=r"C:\Users\fengq\Desktop\EGFR\EGFR的v12"; R=os.path.join(V12,"results"); PK=os.path.join(V12,"dataset_package")
OUT=os.path.join(V12,"03_tables"); os.makedirs(OUT,exist_ok=True)
def rd(p):
    with open(p,encoding="utf-8-sig",newline="") as f: return list(csv.DictReader(f))
def three_line(doc, headers, rows, font=8.5, widths=None):
    t=doc.add_table(rows=1+len(rows), cols=len(headers)); t.style=doc.styles['Normal Table']; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    tp=t._tbl.tblPr
    for old in tp.findall(qn('w:tblBorders')): tp.remove(old)
    b=OxmlElement('w:tblBorders')
    for e in ('top','left','bottom','right','insideH','insideV'):
        el=OxmlElement(f'w:{e}'); el.set(qn('w:val'),'nil'); b.append(el)
    tp.append(b)
    def bd(cell, edges, sz=8):
        tcPr=cell._tc.get_or_add_tcPr()
        for old in tcPr.findall(qn('w:tcBorders')): tcPr.remove(old)
        tb=OxmlElement('w:tcBorders')
        for e in ('top','left','bottom','right','insideH','insideV'):
            el=OxmlElement(f'w:{e}')
            if e in edges: el.set(qn('w:val'),'single'); el.set(qn('w:sz'),str(sz)); el.set(qn('w:color'),'000000')
            else: el.set(qn('w:val'),'nil')
            tb.append(el)
        tcPr.append(tb)
    for j,h in enumerate(headers):
        c=t.cell(0,j); c.text=''; r=c.paragraphs[0].add_run(str(h)); r.bold=True; r.font.size=Pt(font); r.font.name='Times New Roman'
        bd(c,{'top','bottom'},12)
    for i,row in enumerate(rows,1):
        for j,v in enumerate(row):
            c=t.cell(i,j); c.text=''; r=c.paragraphs[0].add_run(str(v)); r.font.size=Pt(font); r.font.name='Times New Roman'
            if i==len(rows): bd(c,{'bottom'},12)
    doc.add_paragraph('')
def heading(doc, text, level=1):
    h=doc.add_heading(level=level); r=h.add_run(text); r.font.name='Times New Roman'; r.font.color.rgb=RGBColor(0,0,0)
# ---------------- Tables ----------------
reg=rd(os.path.join(PK,"01_cohort_registry","cohort_registry.csv")); cc=[r for r in reg if r["design"]=="case-control"]
ctxmap={"GSE32863":"LUAD","GSE19804":"LUAD","GSE10072":"LUAD","GSE44076":"CRC","GSE41258":"CRC","GSE23878":"CRC",
        "GSE75214":"IBD","GSE87473":"IBD","GSE179285":"IBD","GSE4302":"Asthma","GSE43696":"Asthma","GSE67472":"Asthma",
        "GSE27342":"STAD","GSE63089":"STAD","GSE13911":"STAD","GSE15471":"PAAD","GSE28735":"PAAD","GSE62452":"PAAD",
        "GSE57957":"HCC","GSE62232":"HCC","GSE23400":"ESCC","GSE20347":"ESCC","GSE76925":"COPD","GSE47460":"COPD",
        "GSE33814":"NAFLD","GSE66676":"NAFLD"}
order=["LUAD","CRC","STAD","PAAD","HCC","ESCC","IBD","COPD","NAFLD","Asthma"]
agg={}
for r in cc:
    c=ctxmap.get(r["accession"],"other"); a=agg.setdefault(c, dict(cohorts=0, assays=0, pats=0, paired=0, platforms=set(), nopid=0))
    a["cohorts"]+=1; a["assays"]+=int(r["n_assay_records"])
    a["pats"]+=int(r["n_unique_patients"]) if str(r["n_unique_patients"]).isdigit() else 0
    a["paired"]+=int(r["patients_with_both_arms"] or 0) if r["paired_design_used"]=="yes" else 0
    a["platforms"].add(r["platform"]); a["nopid"]+= 1 if r["patient_identity_available"]=="no" else 0
t1_rows=[[c, agg[c]["cohorts"], agg[c]["assays"], agg[c]["pats"], agg[c]["paired"], len(agg[c]["platforms"]), agg[c]["nopid"]] for c in order if c in agg]
d=Document(); st=d.styles['Normal']; st.font.name='Times New Roman'; st.font.size=Pt(10)
heading(d,"Table 1. Cohort composition of the resource by disease context",1)
three_line(d, ["Context","Cohorts","Assay records","Patients (with identifiers)","Paired patients","Platforms","Cohorts without patient identifiers"], t1_rows)
d.add_paragraph("Case-control series only (26 cohorts); one treatment-response series (GSE16879, 43 samples) was reviewed and excluded, see Supplementary Table S4.")
s1=[[r["accession"], ctxmap.get(r["accession"],"IBD" if r["accession"]=="GSE16879" else ""), r["platform"], r["n_assay_records"],
     r["n_unique_patients"], r["n_case"], r["n_control"], r["patients_with_both_arms"], r["paired_design_used"],
     r["patient_identity_available"], r["design"], r["origin_pmid"]] for r in
     sorted(reg, key=lambda x: (order.index(ctxmap.get(x["accession"],"IBD")) if ctxmap.get(x["accession"],"IBD") in order else 99, x["accession"]))]
heading(d,"Supplementary Table S1. Cohort registry with design, counts and origin publications",1)
three_line(d, ["Accession","Context","Platform","Assays","Patients","Cases","Controls","Patients with both arms","Paired design used","Patient IDs available","Design","Origin PMID"], s1, font=7.5)
cov=rd(os.path.join(R,"v12_gene_coverage.csv"))
mods={}
for r in cov:
    m=r["module"]; c=float(r["coverage"]); d0=mods.setdefault(m, dict(members=r["n_members"], minc=1.0, n_inc=0))
    d0["minc"]=min(d0["minc"], c); d0["n_inc"]+= 1 if int(r["n_present"])<int(r["n_members"]) else 0
s2=[[m, v["members"], f"{v['minc']:.2f}", v["n_inc"]] for m,v in sorted(mods.items(), key=lambda kv: kv[1]["minc"])]
heading(d,"Supplementary Table S2. Module membership and cross-cohort gene coverage",1)
three_line(d, ["Module","Members","Minimum coverage across cohorts","Cohorts with incomplete coverage"], s2)
s3=[[r["context"], r["cells_total"], r["cells_used"], r["donors"], r["cell_types"], r["arms"], r["control_kinds"]] for r in rd(os.path.join(R,"v12_sc_dataset_audit.csv"))]
heading(d,"Supplementary Table S3. Single-cell datasets, cells and comparison arms",1)
three_line(d, ["Dataset","Cells total","Cells used","Donors","Cell types","Arms","Control definitions / notes"], s3, font=8)
s4=[[r["item"], r["n"], r["layer"], r["reason"]] for r in rd(os.path.join(PK,"08_qc","exclusion_and_flag_log.csv"))]
heading(d,"Supplementary Table S4. Documented exclusions, flags and design corrections",1)
three_line(d, ["Item","n","Layer","Reason"], s4, font=8)
s5=[["Per-cohort reproducibility",">=2/3 of cohorts share the majority sign and >=2 cohorts at FDR<0.05 (both cohorts when k=2)"],
    ["Common-denominator effect size","unpaired: SMDH; paired: SMCRPH with the empirical within-pair correlation; both use sqrt((sd1^2+sd2^2)/2)"],
    ["Paired design","derived from metadata: >=4 patients contributing both a case and a control sample"],
    ["Meta-analysis","REML with Hartung-Knapp adjustment; DerSimonian-Laird, ad hoc Knapp-Hartung and unpaired-only as sensitivity"],
    ["Composition adjustment","base model (module ~ disease [+ patient]) and adjusted model (+ epi, fib, end, imm) on identical samples; median adjusted/base coefficient ratio reported"],
    ["Composition collinearity","R2 of a linear model of disease status on the four compartment scores, per cohort (median 0.33, max 0.89)"],
    ["Common-gene sensitivity","re-scoring with the members present in every cohort; flagged when effect direction disagrees with full-set scoring in >=20% of cohorts, or when <3 members are shared (not scorable)"],
    ["Proportional hazards","cox.zph on the univariable model with age; violation flagged at p<0.05 (0 of 17 features flagged)"],
    ["Single-cell comparisons","paired signed-rank when >=5 donors contribute both arms; unpaired only when the arms are donor-disjoint; otherwise reported as not testable"],
    ["Mutation status","reported only for samples in the source mutation-profiled list: 'mutation reported' or 'no mutation reported in the selected mutation data'; others labelled not assayed/unknown"]]
heading(d,"Supplementary Table S5. Decision rules and thresholds applied in validation",1)
three_line(d, ["Rule","Definition / threshold"], s5, font=8)
tp=os.path.join(OUT,"Tables_v12.docx"); d.save(tp); print("tables docx:",tp)
# ---------------- Figure legends (separate file) ----------------
L=Document(); st2=L.styles['Normal']; st2.font.name='Times New Roman'; st2.font.size=Pt(11)
heading(L,"Figure legends (v12)",1)
for t,b in [
 ("Figure 1. Resource overview.",
  "(A) Module gene coverage per disease context: the fraction of frozen module members present in the processed expression matrix of each cohort; the colour scale runs from 0.33 (dark red, lowest observed coverage) to 1.00 (pale blue). "
  "(B) Cohort composition by context, shown as three facets: assay records, unique patients with available identifiers, and paired patients (patients contributing both a case and a control sample). "
  "(C) Single-cell datasets with the comparison arms available in each; the lung dataset is shown as excluded because all of its cells come from tumour tissue and no control arm exists."),
 ("Figure 2. Technical validation.",
  "(A) Documented exclusions, flags and non-testable comparisons on a logarithmic scale: adenocarcinoma samples excluded from the squamous oesophageal layer (n=89), in utero cells excluded from the gastric single-cell dataset (n=1,277), tumour-only cells excluded with the lung dataset (n=117,266), malignant-labelled cells found inside non-tumour colorectal samples and flagged rather than reassigned (n=358), and single-cell comparisons that cannot be tested because donors overlap between arms without forming five complete pairs (n=80). "
  "(B) The comparison design actually used in the single-cell layer: paired signed-rank tests on complete donor pairs (red), donor-disjoint unpaired tests (blue), explicitly labelled cell-identity contrasts (pale blue) and comparisons reported as not testable (grey). "
  "(C) The complete estimates layer: the distribution of all per-cohort standardized effects, faceted by the measure used (SMCRPH for paired cohorts with the empirical within-pair correlation; SMDH for unpaired cohorts); non-significant results are retained. "
  "(D) Common-gene sensitivity of module scoring: Spearman correlation between full-member scoring and scoring restricted to the members present in every cohort; red bars mark modules whose effect direction disagreed in at least 20% of cohorts, grey bars mark modules with too few shared members to re-score ('n.a.'), and the dashed line marks the median. "
  "(E) Collinearity between disease status and tissue composition per cohort, measured as the R2 of a linear model of disease status on the four compartment scores, with the median marked; composition adjustment removes part of the disease contrast by construction in cohorts where this R2 is high."),
 ("Supplementary figures.",
  "Supplementary Figure S1. Marker-proxy versus xCell sensitivity of composition adjustment across all states described in the previous version of this resource. "
  "Supplementary Figure S2. Two-step xCell residualization sensitivity, retained for comparison with the base-versus-adjusted models used here.")]:
    heading(L,t,2); L.add_paragraph(b)
lp=os.path.join(OUT,"Figure_legends_v12.docx"); L.save(lp); print("legends docx:",lp)
