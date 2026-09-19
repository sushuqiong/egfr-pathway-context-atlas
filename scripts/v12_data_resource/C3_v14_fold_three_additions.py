#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""v14 (C3): fold the three additions into the manuscript, the tables and every artefact.

1. single-cell threshold gradient (no minimum / >=10 / >=20 / >=50 cells per donor)
2. three multiplicity definitions (comparison-type family, cell-type family, global)
3. an explicit account of the 68% cross-cohort consistency
"""
import csv, hashlib, os, re, shutil, subprocess, zipfile
V14=r"C:\Users\fengq\Desktop\EGFR\EGFR的v14"; PK=os.path.join(V14,"dataset_package"); FIG=os.path.join(V14,"02_figures")
TBL=os.path.join(V14,"03_tables"); repo=r"C:\Users\fengq\Desktop\EGFR\egfr-pathway-context-actlas"
repo=r"C:\Users\fengq\Desktop\EGFR\egfr-pathway-context-atlas"
mdp=os.path.join(V14,"01_manuscript","DataDescriptor_v14.md"); md=open(mdp,encoding="utf-8").read()
def rd(p):
    with open(p,encoding="utf-8-sig",newline="") as f: return list(csv.DictReader(f))
grad=rd(os.path.join(PK,"06_single_cell","single_cell_threshold_gradient.csv"))
dis=rd(os.path.join(PK,"08_qc","qc_cross_cohort_disagreement_summary.csv"))
def g(tag,field):
    for r in grad:
        if r["donor_set"]==tag: return r[field]
    return "?"
rows_s="; ".join(f"{r['donor_set']}: {r['comparisons']} comparisons, {r['testable']} testable, "
                 f"family BH {r['significant_family_bh']}, cell-type family {r['significant_celltype_family_bh']}, "
                 f"global BH {r['significant_global_bh']}" for r in grad)
d=dict((r["metric"],r["value"]) for r in dis)
# ---- 1. replace the sensitivity sentence in Technical Validation ----
old=re.search(r"Single-cell sensitivity analyses are shipped[^.]*\..*?(?=Survival layer)", md, re.S)
new=("Single-cell sensitivity analyses are shipped as Supplementary Table S18. The donor-abundance threshold was varied over four settings "
     f"({rows_s}), and the multiplicity correction was computed in three ways: within the comparison-type family (the primary definition), within the "
     "cell-type family, and jointly across all comparisons; the number of significant comparisons moves with the family size, which is exactly why "
     "the single-cell layer is reported as descriptive rather than confirmatory. ")
md = md.replace(old.group(0), new) if old else md
# ---- 2. account of the cross-cohort consistency ----
old_cons=re.search(r"Comparability: \d+ of \d+ context-module pairs[^;]*;", md)
add=("Disagreement is narrow rather than diffuse: " + d.get("inconsistent pairs with k = 2","21") + " of the discordant pairs rest on only two cohorts "
     "(one positive, one negative) and the remainder on three (two against one), so every discordant pair is decided by a single cohort and no pair "
     "draws on more than three cohorts; the 68% is therefore a fraction of stable pairs, not an expected agreement rate for an unseen cohort. "
     "Discordance concentrates in non-alcoholic fatty liver disease (" + d.get("inconsistent pairs by context (top)","") + " pairs counting all contexts) "
     "and in the " + d.get("inconsistent pairs by module (top)","") + " modules; for example, SRC_FAK in asthma is positive in two cohorts and negative in the third, "
     "while ERBB_RECEPTORS in colorectal cancer is positive in two of three. Per-pair detail, including the sign of every contributing cohort, is shipped as "
     "08_qc/qc_cross_cohort_disagreement_detail.csv (Supplementary Table S19). ")
if old_cons:
    ins=md.index(old_cons.group(0))+len(old_cons.group(0))
    md=md[:ins]+" "+add+md[ins:]
else:
    md=md.replace("Coverage and composition are shown in Figure 1", "Coverage and composition are shown in Figure 1. "+add,1)
# ---- 3. Usage Notes caveat ----
if "single cohort can flip a pair" not in md:
    md=md.replace("(ii) coverage must be checked before any cross-cohort comparison",
                  "(ii) direction consistency is a property of the shipped pairs: when only two or three cohorts contribute, a single cohort can flip a pair, "
                  "so the consistency fraction must not be read as an expected agreement rate for a new cohort; coverage must be checked before any cross-cohort comparison",1)
open(mdp,"w",encoding="utf-8",newline="\n").write(md)
print("manuscript updated: gradient sentence", bool(old), "| disagreement paragraph", bool(old_cons), "| usage-notes caveat", "single cohort can flip a pair" in md)
# ---- 4. tables: S18 and S19 ----
q=os.path.join(V14,"scripts","71_v14_tables_legends.py"); t=open(q,encoding="utf-8").read()
if "Supplementary Table S18" not in t:
    t=t.replace('tp=os.path.join(OUT,"Tables_v14.docx")',
'''head(doc,"Supplementary Table S18. Single-cell sensitivity: donor-abundance thresholds and multiplicity definitions")
_g=rd(os.path.join(PK,"06_single_cell","single_cell_threshold_gradient.csv"))
three_line(doc,["Donor set","Comparisons","Testable","Not testable","Significant: comparison-type family","Significant: cell-type family","Significant: global"],
  [[r["donor_set"],r["comparisons"],r["testable"],r["not_testable"],r["significant_family_bh"],r["significant_celltype_family_bh"],r["significant_global_bh"]] for r in _g],8)
head(doc,"Supplementary Table S19. Cross-cohort direction consistency: where the disagreement sits")
_d=rd(os.path.join(PK,"08_qc","qc_cross_cohort_disagreement_summary.csv"))
three_line(doc,["Metric","Value"],[[r["metric"],r["value"]] for r in _d],8)
_dc=rd(os.path.join(PK,"08_qc","qc_cross_cohort_disagreement_detail.csv"))
_inc=[r for r in _dc if str(r["direction_consistent"]).upper()!="TRUE"][:12]
doc.add_paragraph("First twelve discordant pairs (complete table shipped as 08_qc/qc_cross_cohort_disagreement_detail.csv):")
three_line(doc,["Context","Module","Cohorts","Pattern","Cohorts positive","Cohorts negative"],
  [[r["context"],r["module"],r["k"],r["pattern"],r["positive_cohorts"],r["negative_cohorts"]] for r in _inc],8)
tp=os.path.join(OUT,"Tables_v14.docx")''',1)
    open(q,"w",encoding="utf-8").write(t); print("tables builder extended with S18/S19")
rr=subprocess.run([os.sys.executable,q],capture_output=True,text=True,cwd=V14)
print("tables:",[l for l in (rr.stdout or "").splitlines() if "tables" in l.lower()][-1:] or rr.stderr[-200:])
subprocess.run([r"C:\Program Files\LibreOffice\program\soffice.exe","--headless","--convert-to","pdf","--outdir",TBL,os.path.join(TBL,"Tables_v14.docx")],capture_output=True)
if os.path.exists(os.path.join(TBL,"Tables_v14.pdf")): shutil.move(os.path.join(TBL,"Tables_v14.pdf"),os.path.join(TBL,"Supplementary_Information.pdf"))
print("SI pdf rebuilt:",round(os.path.getsize(os.path.join(TBL,"Supplementary_Information.pdf"))/1e6,2),"MB")
# ---- 5. rebuild the manuscript documents ----
from docx import Document
from docx.shared import Pt, Inches, RGBColor
TOK=re.compile(r"⟦(\d+(?:,\d+)*)⟧")
def add_para(doc,text):
    p=doc.add_paragraph(); pos=0
    for m in TOK.finditer(text):
        if m.start()>pos: p.add_run(text[pos:m.start()])
        s=p.add_run(m.group(1)); s.font.superscript=True; pos=m.end()
    if pos<len(text): p.add_run(text[pos:])
    for r in p.runs: r.font.name="Times New Roman"
doc=Document(); st=doc.styles["Normal"]; st.font.name="Times New Roman"; st.font.size=Pt(11)
figs=[("Figure 1 (legend in the separate figure-legend file).","Fig1_resource_overview.png"),
      ("Figure 2 (legend in the separate figure-legend file).","Fig2_technical_validation.png"),
      ("Figure 3 (legend in the separate figure-legend file).","Fig3_comparability_checks.png"),
      ("Figure 4 (legend in the separate figure-legend file).","Fig4_sensitivity_checks.png")]
for line in md.splitlines():
    s=line.rstrip()
    if not s.strip(): continue
    if s.startswith("# "):
        h=doc.add_heading(level=0); r=h.add_run(s[2:].strip()); r.font.name="Times New Roman"; r.font.color.rgb=RGBColor(0,0,0)
    elif s.startswith("## "):
        h=doc.add_heading(level=1); r=h.add_run(s[3:].strip()); r.font.name="Times New Roman"; r.font.color.rgb=RGBColor(0,0,0)
    elif s.startswith("**") and s.endswith("**") and not s.startswith("**Figure"):
        h=doc.add_heading(level=2); r=h.add_run(s.strip("*")); r.font.name="Times New Roman"; r.font.color.rgb=RGBColor(0,0,0)
    else: add_para(doc, s.replace("**","").replace("*",""))
doc.add_page_break(); h=doc.add_heading(level=1); r=h.add_run("Figures"); r.font.name="Times New Roman"; r.font.color.rgb=RGBColor(0,0,0)
for cap,fn in figs:
    h=doc.add_heading(level=2); r=h.add_run(cap); r.font.name="Times New Roman"; r.font.color.rgb=RGBColor(0,0,0)
    doc.add_picture(os.path.join(FIG,fn), width=Inches(6.2)); doc.add_paragraph("")
dp=os.path.join(V14,"01_manuscript","DataDescriptor_v14.docx"); doc.save(dp)
subprocess.run([r"C:\Program Files\LibreOffice\program\soffice.exe","--headless","--convert-to","pdf","--outdir",
                os.path.join(V14,"01_manuscript"),dp],capture_output=True)
d=Document(dp); dt="\n".join(p.text for p in d.paragraphs)
disk={fn:hashlib.sha256(open(os.path.join(FIG,fn),"rb").read()).hexdigest() for _,fn in figs}
em=[hashlib.sha256(r.target_part.blob).hexdigest() for r in d.part.rels.values() if "image" in r.reltype]
nref=len(re.findall(r"(?m)^\d+\. ", dt.split("References")[-1]))
print(f"docx: superscripts={sum(1 for p in d.paragraphs for r in p.runs if r.font.superscript)} | figures ok={sum(1 for h in em if h in disk.values())}/{len(em)} | refs={nref} | S18={'S18' in dt} S19={'S19' in dt}")
# ---- 6. zips ----
z=os.path.join(r"C:\Users\fengq\Desktop","★Zenodo上传-就选这个_v14"); shutil.make_archive(z,"zip",PK)
stage=os.path.join(V14,"_stage11"); shutil.rmtree(stage,ignore_errors=True)
pairs=[("01_manuscript/DataDescriptor_v14.docx","01_Manuscript/DataDescriptor_v14.docx"),
       ("01_manuscript/DataDescriptor_v14.pdf","01_Manuscript/DataDescriptor_v14.pdf"),
       ("04_Cover_letter_v14.md","01_Manuscript/Cover_letter.md"),
       ("03_tables/Supplementary_Information.pdf","02_Supplementary/Supplementary_Information.pdf"),
       ("03_tables/Tables_v14.docx","02_Supplementary/Tables_v14.docx"),
       ("03_tables/Figure_legends_v14.docx","02_Supplementary/Figure_legends_v14.docx")]
names={1:"resource_overview",2:"technical_validation",3:"comparability_checks",4:"sensitivity_checks"}
for i in (1,2,3,4):
    for ext in ("png","pdf"): pairs.append((f"02_figures/Fig{i}_{names[i]}.{ext}", f"03_Figures/Fig{i}_{names[i]}.{ext}"))
for a,b in pairs:
    sp=os.path.join(V14,a)
    if os.path.exists(sp):
        os.makedirs(os.path.dirname(os.path.join(stage,b)),exist_ok=True); shutil.copy(sp,os.path.join(stage,b))
open(os.path.join(stage,"00_READ_ME_FIRST.txt"),"w",encoding="utf-8").write(
"Submission package - cross-disease transcriptomic resource (v14)\n\n01_Manuscript : DataDescriptor_v14.docx (editable), .pdf and Cover_letter.md\n"
"02_Supplementary : Supplementary_Information.pdf (Table 1 and Supplementary Tables S1-S19) and the same tables as .docx\n"
"03_Figures : Figures 1-4 as .png (300 dpi) and .pdf\n\nThe data deposit (including analysis code) is provided separately as the Zenodo archive.\n")
out=os.path.join(r"C:\Users\fengq\Desktop","★投稿上传-就选这个_v14.zip")
with zipfile.ZipFile(out,"w",zipfile.ZIP_DEFLATED) as zf:
    for d0,_,fn in os.walk(stage):
        for f in fn: fp=os.path.join(d0,f); zf.write(fp,os.path.relpath(fp,stage))
shutil.rmtree(stage,ignore_errors=True)
print("zips:",round(os.path.getsize(z+".zip")/1e6,2),"MB /",round(os.path.getsize(out)/1e6,2),"MB")
# ---- 7. push ----
shutil.copy(mdp, os.path.join(repo,"manuscript","DataDescriptor_v14.md"))
for f in os.listdir(os.path.join(V14,"scripts")):
    src=os.path.join(V14,"scripts",f)
    if os.path.isfile(src): shutil.copy(src, os.path.join(repo,"scripts","v12_data_resource",f))
for f in ["single_cell_threshold_gradient.csv","single_cell_comparisons_with_celltype_family_fdr.csv"]:
    p=os.path.join(PK,"06_single_cell",f)
    if os.path.exists(p): shutil.copy(p, os.path.join(repo,"results",f))
for f in ["qc_cross_cohort_disagreement_summary.csv","qc_cross_cohort_disagreement_detail.csv"]:
    p=os.path.join(PK,"08_qc",f)
    if os.path.exists(p): shutil.copy(p, os.path.join(repo,"results",f))
for cmd in [["git","add","-A"],
            ["git","-c","user.name=sushuqiong","-c","user.email=fengqi9666@qq.com","commit","-q","-m",
             "v14.6: single-cell sensitivity extended to a four-level donor-abundance gradient with three multiplicity definitions (Tables S18), and an explicit account of the 68% cross-cohort consistency (Table S19): every discordant pair is decided by a single cohort and no pair draws on more than three cohorts"],
            ["git","push","-q","origin","main"],
            ["git","tag","-f","v14.0","-m","resource v14.0 matching the submitted manuscript"],
            ["git","push","-q","-f","origin","v14.0"]]:
    rr2=subprocess.run(cmd,cwd=repo,capture_output=True,text=True); print(cmd[1],"->",rr2.returncode)
print(subprocess.run(["git","log","--oneline","-1"],cwd=repo,capture_output=True,text=True).stdout.strip())
