#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""v14 (B6): rebuild every artefact from the finished manuscript, embed Table 1, add a cover letter,
re-zip both packages, tag the repository release and push."""
import csv, hashlib, os, re, shutil, subprocess, zipfile
V14=r"C:\Users\fengq\Desktop\EGFR\EGFR的v14"; PK=os.path.join(V14,"dataset_package"); TBL=os.path.join(V14,"03_tables")
FIG=os.path.join(V14,"02_figures"); repo=r"C:\Users\fengq\Desktop\EGFR\egfr-pathway-context-atlas"
mdp=os.path.join(V14,"01_manuscript","DataDescriptor_v14.md"); md=open(mdp,encoding="utf-8").read()
def rd(p):
    with open(p,encoding="utf-8-sig",newline="") as f: return list(csv.DictReader(f))
# ---- Table 1 data ----
CTX={"GSE32863":"LUAD","GSE19804":"LUAD","GSE10072":"LUAD","GSE44076":"CRC","GSE41258":"CRC","GSE23878":"CRC",
     "GSE75214":"IBD","GSE87473":"IBD","GSE179285":"IBD","GSE4302":"Asthma","GSE43696":"Asthma","GSE67472":"Asthma",
     "GSE27342":"STAD","GSE63089":"STAD","GSE13911":"STAD","GSE15471":"PAAD","GSE28735":"PAAD","GSE62452":"PAAD",
     "GSE57957":"HCC","GSE62232":"HCC","GSE23400":"ESCC","GSE20347":"ESCC","GSE76925":"COPD","GSE47460":"COPD",
     "GSE33814":"NAFLD","GSE66676":"NAFLD"}
order=["LUAD","CRC","STAD","PAAD","HCC","ESCC","IBD","COPD","NAFLD","Asthma"]
reg=[r for r in rd(os.path.join(PK,"01_cohort_registry","cohort_registry.csv")) if r["design"]=="case-control"]
agg={}
for r in reg:
    c=CTX.get(r["accession"],"other"); a=agg.setdefault(c,dict(n=0,assays=0,pat=0,pair=0,plat=set()))
    a["n"]+=1; a["assays"]+=int(r["n_assay_records"])
    a["pat"]+=int(r["n_unique_patients"]) if str(r["n_unique_patients"]).isdigit() else 0
    a["pair"]+=int(r["patients_with_both_arms"] or 0) if r["paired_design_used"]=="yes" else 0
    a["plat"].add(r["platform"])
T1=[[c,agg[c]["n"],agg[c]["assays"],agg[c]["pat"],agg[c]["pair"],len(agg[c]["plat"])] for c in order if c in agg]
# ---- docx builder ----
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
TOK=re.compile(r"⟦(\d+(?:,\d+)*)⟧")
def add_para(doc,text):
    p=doc.add_paragraph(); pos=0
    for m in TOK.finditer(text):
        if m.start()>pos: p.add_run(text[pos:m.start()])
        s=p.add_run(m.group(1)); s.font.superscript=True; pos=m.end()
    if pos<len(text): p.add_run(text[pos:])
    for r in p.runs: r.font.name="Times New Roman"
def three_line(doc,headers,rows,font=8.5):
    t=doc.add_table(rows=1+len(rows),cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    tp=t._tbl.tblPr
    for old in tp.findall(qn('w:tblBorders')): tp.remove(old)
    b=OxmlElement('w:tblBorders')
    for e in ('top','left','bottom','right','insideH','insideV'):
        el=OxmlElement(f'w:{e}'); el.set(qn('w:val'),'nil'); b.append(el)
    tp.append(b)
    def bd(cell,edges,sz=8):
        tcPr=cell._tc.get_or_add_tcPr()
        for old in tcPr.findall(qn('w:tcBorders')): tcPr.remove(old)
        tb=OxmlElement('w:tcBorders')
        for e in ('top','left','bottom','right','insideH','insideV'):
            el=OxmlElement(f'w:{e}')
            if e in edges: el.set(qn('w:val'),'single'); el.set(qn('w:sz'),str(sz)); el.set(qn('w:color'),'000000')
            else: el.set(qn('w:val'),'nil')
            tb.append(el)
        tcPr.append(tb)
    for j,h in enumerate(headers):
        c=t.cell(0,j); c.text=''; rr=c.paragraphs[0].add_run(str(h)); rr.bold=True; rr.font.size=Pt(font); rr.font.name="Times New Roman"
        bd(c,{'top','bottom'},12)
    for i,row in enumerate(rows,1):
        for j,v in enumerate(row):
            c=t.cell(i,j); c.text=''; rr=c.paragraphs[0].add_run(str(v)); rr.font.size=Pt(font); rr.font.name="Times New Roman"
            if i==len(rows): bd(c,{'bottom'},12)
    doc.add_paragraph('')
doc=Document(); st=doc.styles["Normal"]; st.font.name="Times New Roman"; st.font.size=Pt(11)
figs=[("Figure 1 (legend in the separate figure-legend file).","Fig1_resource_overview.png"),
      ("Figure 2 (legend in the separate figure-legend file).","Fig2_technical_validation.png"),
      ("Figure 3 (legend in the separate figure-legend file).","Fig3_comparability_checks.png"),
      ("Figure 4 (legend in the separate figure-legend file).","Fig4_sensitivity_checks.png")]
inserted=False
for line in md.splitlines():
    s=line.rstrip()
    if not s.strip(): continue
    if s.startswith("# "):
        h=doc.add_heading(level=0); r=h.add_run(s[2:].strip()); r.font.name="Times New Roman"; r.font.color.rgb=RGBColor(0,0,0)
    elif s.startswith("## "):
        h=doc.add_heading(level=1); r=h.add_run(s[3:].strip()); r.font.name="Times New Roman"; r.font.color.rgb=RGBColor(0,0,0)
    elif s.startswith("**") and s.endswith("**") and not s.startswith("**Figure"):
        h=doc.add_heading(level=2); r=h.add_run(s.strip("*")); r.font.name="Times New Roman"; r.font.color.rgb=RGBColor(0,0,0)
        if s.startswith("**Technical Validation**") and not inserted:
            hh=doc.add_heading(level=2); rr=hh.add_run("Table 1. Cohort composition of the resource by disease context")
            rr.font.name="Times New Roman"; rr.font.color.rgb=RGBColor(0,0,0)
            three_line(doc,["Context","Cohorts","Assay records","Patients with identifiers","Patients with both tissues","Platforms"],T1)
            inserted=True
    else: add_para(doc, s.replace("**","").replace("*",""))
doc.add_page_break(); h=doc.add_heading(level=1); r=h.add_run("Figures"); r.font.name="Times New Roman"; r.font.color.rgb=RGBColor(0,0,0)
for cap,fn in figs:
    h=doc.add_heading(level=2); r=h.add_run(cap); r.font.name="Times New Roman"; r.font.color.rgb=RGBColor(0,0,0)
    doc.add_picture(os.path.join(FIG,fn), width=Inches(6.2)); doc.add_paragraph("")
dp=os.path.join(V14,"01_manuscript","DataDescriptor_v14.docx"); doc.save(dp)
subprocess.run([r"C:\Program Files\LibreOffice\program\soffice.exe","--headless","--convert-to","pdf","--outdir",
                os.path.join(V14,"01_manuscript"),dp],capture_output=True)
d=Document(dp); dt="\n".join(p.text for p in d.paragraphs)
disk={fn:hashlib.sha256(open(os.path.join(FIG,fn),"rb").read()).hexdigest() for _,fn in figs}
em=[hashlib.sha256(r.target_part.blob).hexdigest() for r in d.part.rels.values() if "image" in r.reltype]
n_refs_listed=len(re.findall(r"(?m)^\d+\. ", dt.split("References")[-1]))
print(f"docx: tables={len(d.tables)} | superscripts={sum(1 for p in d.paragraphs for r in p.runs if r.font.superscript)} | "
      f"figures matching disk={sum(1 for h in em if h in disk.values())}/{len(em)} | refs listed={n_refs_listed}")
# ---- cover letter ----
open(os.path.join(V14,"04_Cover_letter_v14.md"),"w",encoding="utf-8").write(
"Dear Editors of Scientific Data,\n\n"
"We submit for consideration as a Data Descriptor the manuscript \"Coverage-aware harmonized cross-disease transcriptome resource "
"with expression-derived compartment scores\".\n\n"
"The resource curates 26 public case-control bulk transcriptome cohorts across ten disease contexts (2,711 assay records, 1,086 patients "
"with identifiers, 12 platforms), together with six tumour molecular contexts, four single-cell datasets and an archived analysis and "
"quality-control pipeline. Its distinguishing feature is that comparability is quantified rather than assumed: every effect row carries the "
"actual gene coverage of its module, cross-cohort direction consistency is reported per context and module, and non-significant, not-estimable "
"and not-testable results are retained in the shipped tables. The deposit contains the data tables, a machine-readable inventory with SHA-256 "
"checksums, runnable reuse examples and the analysis code.\n\n"
"All data are public and de-identified; no ethical approval was required. The resource is deposited in Zenodo under CC-BY-4.0 and the code is "
"available in a public repository tagged for this submission. The manuscript is not under consideration elsewhere.\n\n"
"Sincerely,\nShuqiong Su\nDepartment of Gastroenterology, Guangxi Medical University Cancer Hospital, Nanning, China\nliuaiqun_2004@163.com\n")
# ---- zips ----
z=os.path.join(r"C:\Users\fengq\Desktop","★Zenodo上传-就选这个_v14"); shutil.make_archive(z,"zip",PK)
stage=os.path.join(V14,"_stage10"); shutil.rmtree(stage,ignore_errors=True)
pairs=[("01_manuscript/DataDescriptor_v14.docx","01_Manuscript/DataDescriptor_v14.docx"),
       ("01_manuscript/DataDescriptor_v14.pdf","01_Manuscript/DataDescriptor_v14.pdf"),
       ("04_Cover_letter_v14.md","01_Manuscript/Cover_letter.md"),
       ("03_tables/Supplementary_Information.pdf","02_Supplementary/Supplementary_Information.pdf"),
       ("03_tables/Tables_v14.docx","02_Supplementary/Tables_v14.docx"),
       ("03_tables/Figure_legends_v14.docx","02_Supplementary/Figure_legends_v14.docx")]
names={1:"resource_overview",2:"technical_validation",3:"comparability_checks",4:"sensitivity_checks"}
for i in (1,2,3,4):
    for ext in ("png","pdf"): pairs.append((f"02_figures/Fig{i}_{names[i]}.{ext}", f"03_Figures/Fig{i}_{names[i]}.{ext}"))
for a,b in pairs:
    sp=os.path.join(V14,a)
    if os.path.exists(sp):
        os.makedirs(os.path.dirname(os.path.join(stage,b)),exist_ok=True); shutil.copy(sp,os.path.join(stage,b))
open(os.path.join(stage,"00_READ_ME_FIRST.txt"),"w",encoding="utf-8").write(
"Submission package - cross-disease transcriptomic resource (v14)\n\n01_Manuscript : DataDescriptor_v14.docx (editable), .pdf and Cover_letter.md\n"
"02_Supplementary : Supplementary_Information.pdf (Table 1 and Supplementary Tables S1-S17) and the same tables as .docx\n"
"03_Figures : Figures 1-4 as .png (300 dpi) and .pdf\n\nThe data deposit (including analysis code) is provided separately as the\n"
"Zenodo archive.\n")
out=os.path.join(r"C:\Users\fengq\Desktop","★投稿上传-就选这个_v14.zip")
with zipfile.ZipFile(out,"w",zipfile.ZIP_DEFLATED) as zf:
    for d0,_,fn in os.walk(stage):
        for f in fn: fp=os.path.join(d0,f); zf.write(fp,os.path.relpath(fp,stage))
shutil.rmtree(stage,ignore_errors=True)
print("zips:",round(os.path.getsize(z+".zip")/1e6,2),"MB /",round(os.path.getsize(out)/1e6,2),"MB")
# ---- git: copy, tag, push ----
shutil.copy(mdp, os.path.join(repo,"manuscript","DataDescriptor_v14.md"))
for f in os.listdir(os.path.join(V14,"scripts")):
    src=os.path.join(V14,"scripts",f)
    if os.path.isfile(src): shutil.copy(src, os.path.join(repo,"scripts","v12_data_resource",f))
for cmd in [["git","add","-A"],
            ["git","-c","user.name=sushuqiong","-c","user.email=fengqi9666@qq.com","commit","-q","-m",
             "v14.5: references in Scientific Data/Nature style (52 entries, all with resolvable links, data citations for datasets); editor-pass fixes - full accession list in Data Availability, code availability with release tag, author/keyword block, Table 1 embedded in the manuscript, cover letter added, Technical Validation restricted to data-quality claims"],
            ["git","push","-q","origin","main"],
            ["git","tag","-f","v14.0","-m","resource v14.0 matching the submitted manuscript"],
            ["git","push","-q","-f","origin","v14.0"]]:
    rr=subprocess.run(cmd,cwd=repo,capture_output=True,text=True); print(cmd[1],"->",rr.returncode, (rr.stderr or "")[:80])
print(subprocess.run(["git","log","--oneline","-1"],cwd=repo,capture_output=True,text=True).stdout.strip())
print("tag:", subprocess.run(["git","describe","--tags","--always"],cwd=repo,capture_output=True,text=True).stdout.strip())
