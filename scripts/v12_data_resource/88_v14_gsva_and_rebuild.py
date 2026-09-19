#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""v14: add the GSVA citation where the section actually uses it, then rebuild the manuscript files and ZIPs."""
import os, subprocess, shutil, zipfile
V14=r"C:\Users\fengq\Desktop\EGFR\EGFR的v14"; PK=os.path.join(V14,"dataset_package")
mdp=os.path.join(V14,"01_manuscript","DataDescriptor_v14.md"); md=open(mdp,encoding="utf-8").read()
if "GSVA7" not in md:
    md=md.replace("bulk GSVA z-scores","bulk GSVA7 z-scores",1)
    md=md.replace("it is a GSVA value","it is a GSVA7 value",1)
    open(mdp,"w",encoding="utf-8").write(md)
print("GSVA cited:", "GSVA7" in open(mdp,encoding="utf-8").read())
from docx import Document
from docx.shared import Pt, Inches, RGBColor
md=open(mdp,encoding="utf-8").read()
d=Document(); st=d.styles["Normal"]; st.font.name="Times New Roman"; st.font.size=Pt(11)
def H(x,l):
    h=d.add_heading(level=l); r=h.add_run(x); r.font.name="Times New Roman"; r.font.color.rgb=RGBColor(0,0,0)
for line in md.splitlines():
    s=line.rstrip()
    if not s.strip(): continue
    if s.startswith("# "): H(s[2:].strip(),0)
    elif s.startswith("## "): H(s[3:].strip(),1)
    elif s.startswith("**") and s.endswith("**"): H(s.strip("*"),2)
    else: d.add_paragraph(s.replace("**","").replace("*",""))
d.add_page_break(); H("Figures",1)
for cap,fn in [("Figure 1 (caption above and in the separate legend file).","Fig1_resource_overview.png"),
               ("Figure 2 (caption above and in the separate legend file).","Fig2_technical_validation.png"),
               ("Figure 3 (caption above and in the separate legend file).","Fig3_comparability_checks.png"),
               ("Figure 4 (caption above and in the separate legend file).","Fig4_sensitivity_checks.png")]:
    H(cap,2); pp=os.path.join(V14,"02_figures",fn)
    if os.path.exists(pp): d.add_picture(pp,width=Inches(6.2))
    d.add_paragraph("")
dp=os.path.join(V14,"01_manuscript","DataDescriptor_v14.docx"); d.save(dp)
subprocess.run([r"C:\Program Files\LibreOffice\program\soffice.exe","--headless","--convert-to","pdf","--outdir",
                os.path.join(V14,"01_manuscript"),dp],capture_output=True)
z=os.path.join(r"C:\Users\fengq\Desktop","★Zenodo上传-就选这个_v14"); shutil.make_archive(z,"zip",PK)
stage=os.path.join(V14,"_stage4"); shutil.rmtree(stage,ignore_errors=True); os.makedirs(stage)
for a,b in [("01_manuscript/DataDescriptor_v14.docx","01_Manuscript/DataDescriptor_v14.docx"),
            ("01_manuscript/DataDescriptor_v14.pdf","01_Manuscript/DataDescriptor_v14.pdf"),
            ("03_tables/Tables_v14.docx","02_Tables/Tables_v14.docx"),
            ("03_tables/Figure_legends_v14.docx","02_Tables/Figure_legends_v14.docx"),
            ("02_figures/Fig1_resource_overview.png","03_Figures/Fig1_resource_overview.png"),
            ("02_figures/Fig2_technical_validation.png","03_Figures/Fig2_technical_validation.png"),
            ("02_figures/Fig3_comparability_checks.png","03_Figures/Fig3_comparability_checks.png"),
            ("02_figures/Fig4_sensitivity_checks.png","03_Figures/Fig4_sensitivity_checks.png"),
            ("02_figures/Fig1_resource_overview.pdf","03_Figures/Fig1_resource_overview.pdf"),
            ("02_figures/Fig2_technical_validation.pdf","03_Figures/Fig2_technical_validation.pdf"),
            ("02_figures/Fig3_comparability_checks.pdf","03_Figures/Fig3_comparability_checks.pdf"),
            ("02_figures/Fig4_sensitivity_checks.pdf","03_Figures/Fig4_sensitivity_checks.pdf"),
            ("06_三路冷读审稿意见与处置_v14.md","00_先看这个_三路冷读处置.md"),
            ("05_审稿意见与处置_v14.md","00_先看这个_GPT6意见处置.md"),
            ("03_Zenodo上传指南.md","00_先看这个_Zenodo上传指南.md")]:
    sp=os.path.join(V14,a)
    if os.path.exists(sp):
        os.makedirs(os.path.dirname(os.path.join(stage,b)),exist_ok=True); shutil.copy(sp,os.path.join(stage,b))
out=os.path.join(r"C:\Users\fengq\Desktop","★投稿上传-就选这个_v14.zip")
with zipfile.ZipFile(out,"w",zipfile.ZIP_DEFLATED) as zf:
    for d0,_,fn in os.walk(stage):
        for f in fn: fp=os.path.join(d0,f); zf.write(fp,os.path.relpath(fp,stage))
shutil.rmtree(stage,ignore_errors=True)
repo=r"C:\Users\fengq\Desktop\EGFR\egfr-pathway-context-atlas"
shutil.copy(mdp, os.path.join(repo,"manuscript","DataDescriptor_v14.md"))
for cmd in [["git","add","-A"],["git","-c","user.name=sushuqiong","-c","user.email=fengqi9666@qq.com","commit","-q","-m",
             "v14.1: GSVA citation added; manuscript and ZIPs rebuilt"],
            ["git","push","-q","origin","main"]]:
    r=subprocess.run(cmd,cwd=repo,capture_output=True,text=True); print(cmd[1],"->",r.returncode)
print("zips:",round(os.path.getsize(z+".zip")/1e6,2),"MB /",round(os.path.getsize(out)/1e6,2),"MB")
print(subprocess.run(["git","log","--oneline","-1"],cwd=repo,capture_output=True,text=True).stdout.strip())
