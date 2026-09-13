#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
V11=r"C:\Users\fengq\Desktop\EGFR\EGFR的v11"
mdp=os.path.join(V11,"01_稿件正文","Manuscript_draft_v11.md"); P=os.path.join(V11,"02_图片")
def three(doc,headers,rows,font=9):
    t=doc.add_table(rows=1+len(rows),cols=len(headers)); t.style=doc.styles['Normal Table']; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    tp=t._tbl.tblPr
    for old in tp.findall(qn('w:tblBorders')): tp.remove(old)
    b=OxmlElement('w:tblBorders')
    for e in ('top','left','bottom','right','insideH','insideV'):
        el=OxmlElement(f'w:{e}'); el.set(qn('w:val'),'nil'); b.append(el)
    tp.append(b)
    def bd(cell,edges,sz):
        tcPr=cell._tc.get_or_add_tcPr()
        for old in tcPr.findall(qn('w:tcBorders')): tcPr.remove(old)
        tb=OxmlElement('w:tcBorders')
        for e in ('top','left','bottom','right','insideH','insideV'):
            el=OxmlElement(f'w:{e}')
            if e in edges: el.set(qn('w:val'),'single'); el.set(qn('w:sz'),str(sz)); el.set(qn('w:color'),'000000')
            else: el.set(qn('w:val'),'nil')
            tb.append(el)
        tcPr.append(tb)
    for j,h in enumerate(headers):
        c=t.cell(0,j); c.text=''; r=c.paragraphs[0].add_run(h); r.bold=True; r.font.size=Pt(font); r.font.name='Times New Roman'; bd(c,{'top','bottom'},12)
    for i,row in enumerate(rows,1):
        for j,v in enumerate(row):
            c=t.cell(i,j); c.text=''; r=c.paragraphs[0].add_run(str(v)); r.font.size=Pt(font); r.font.name='Times New Roman'
            if i==len(rows): bd(c,{'bottom'},12)
    doc.add_paragraph('')
doc=Document(); st=doc.styles['Normal']; st.font.name='Times New Roman'; st.font.size=Pt(11)
def H(x,l):
    h=doc.add_heading(level=l); r=h.add_run(x); r.font.name='Times New Roman'; r.font.color.rgb=RGBColor(0,0,0)
lines=open(mdp,encoding='utf-8').read().splitlines(); skip=False
for line in lines:
    s=line.rstrip()
    if not s.strip(): continue
    if s.startswith('## Tables'):
        H('Tables',1)
        three(doc,['Context','Per-cohort robust','Meta REML/Knha (BH)','Meta DL (BH)','Composition-robust','xCell-resid (sensitivity)','External OS'],
          [['LUAD','14','2','12','0','0','—'],
           ['CRC','11','11','12','10','10','VEGF direction only (not stage-robust)'],
           ['STAD','10','0','9','9','8','3/5 direction only (ACRG)'],
           ['PAAD','10','1','8','0','0','none replicated (GSE21501)'],
           ['HCC','0','0','6','0','0','—'],
           ['ESCC','4','0','8','1','1','—'],
           ['IBD','6','3','10','1','0','—'],
           ['COPD','0','0','5','0','0','—'],
           ['NAFLD','0','0','0','0','0','—'],
           ['Asthma','0','0','4','0','0','—']])
        skip=True; continue
    if skip:
        if s.startswith(('## ','# ')): skip=False
        else: continue
    if s.startswith('# '): H(s[2:].strip(),0); continue
    if s.startswith('## '): H(s[3:].strip(),1); continue
    if s.startswith('### '): H(s[4:].strip(),2); continue
    doc.add_paragraph(s.strip())
fb=[('Figure 1. Parallel evidence layers (not a sequential filter).',[('','fig1_parallel_layers_v11.png')]),
    ('Figure 2. Module states. (A) raw; (B) adjusted model from the base-versus-adjusted comparison.',[('A','fig2a_raw_v11.png'),('B','fig2b_adjusted_v11.png')]),
    ('Figure 3. Base versus adjusted disease coefficients (median adjusted/base ratio 0.87).',[('','fig3_base_vs_adjusted_v11.png')]),
    ('Figure 4. TCGA module-OS associations (FDR<0.05).',[('','fig4_survival_v11.png')]),
    ('Figure 5. Molecular context with profiled-only denominators. (A) amplification; (B) mutation.',[('','fig5_molecular_v11.png')]),
    ('Figure 6. Single-cell comparisons. (A) same-cell-type comparisons with complete pairs annotated; (B) cell-identity contrasts (labelled separately).',[('','fig6_sc_v11.png')]),
    ('Supplementary Figure S1. Marker-proxy versus xCell medians.',[('','fig_paper_xcell_vs_marker.png')]),
    ('Supplementary Figure S2. xCell residualization heatmap (sensitivity).',[('','fig_paper_xcell_heatmap.png')])]
doc.add_page_break(); H('Figures',1)
for cap,pan in fb:
    H(cap,2)
    for tag,fn in pan:
        fp=os.path.join(P,fn)
        if tag: doc.add_paragraph(tag)
        if os.path.exists(fp): doc.add_picture(fp,width=Inches(6.2)); doc.add_paragraph()
        else: doc.add_paragraph('(missing: '+fn+')')
doc.save(os.path.join(V11,"01_稿件正文","Manuscript_draft_v11.docx"))
print("v11 docx written")
