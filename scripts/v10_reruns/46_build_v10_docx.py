#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
V10=r"C:\Users\fengq\Desktop\EGFR\EGFR的v10"
mdp=os.path.join(V10,"01_稿件正文","Manuscript_draft_v10.md"); P=os.path.join(V10,"02_图片")
def three(doc,headers,rows,font=9):
    t=doc.add_table(rows=1+len(rows),cols=len(headers)); t.style=doc.styles['Normal Table']; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    tp=t._tbl.tblPr
    for old in tp.findall(qn('w:tblBorders')): tp.remove(old)
    b=OxmlElement('w:tblBorders')
    for e in ('top','left','bottom','right','insideH','insideV'):
        el=OxmlElement(f'w:{e}'); el.set(qn('w:val'),'nil'); b.append(el)
    tp.append(b)
    def bd(cell,edges,sz):
        tcPr=cell._tc.get_or_add_tcPr()
        for old in tcPr.findall(qn('w:tcBorders')): tcPr.remove(old)
        tb=OxmlElement('w:tcBorders')
        for e in ('top','left','bottom','right','insideH','insideV'):
            el=OxmlElement(f'w:{e}')
            if e in edges: el.set(qn('w:val'),'single'); el.set(qn('w:sz'),str(sz)); el.set(qn('w:color'),'000000')
            else: el.set(qn('w:val'),'nil')
            tb.append(el)
        tcPr.append(tb)
    for j,h in enumerate(headers):
        c=t.cell(0,j); c.text=''; r=c.paragraphs[0].add_run(h); r.bold=True; r.font.size=Pt(font); r.font.name='Times New Roman'; bd(c,{'top','bottom'},12)
    for i,row in enumerate(rows,1):
        for j,v in enumerate(row):
            c=t.cell(i,j); c.text=''; r=c.paragraphs[0].add_run(str(v)); r.font.size=Pt(font); r.font.name='Times New Roman'
            if i==len(rows): bd(c,{'bottom'},12)
    doc.add_paragraph('')
doc=Document(); st=doc.styles['Normal']; st.font.name='Times New Roman'; st.font.size=Pt(11)
def H(x,l):
    h=doc.add_heading(level=l); r=h.add_run(x); r.font.name='Times New Roman'; r.font.color.rgb=RGBColor(0,0,0)
lines=open(mdp,encoding='utf-8').read().splitlines(); skip=False
for line in lines:
    s=line.rstrip()
    if not s.strip(): continue
    if s.startswith('## Tables'):
        H('Tables',1)
        three(doc,['Context','Per-cohort robust','Meta REML/Knha (BH)','Meta DL (BH)','Joint-model robust','xCell-resid (sensitivity)','External OS'],
          [['LUAD','14','0','0','0','0','—'],
           ['CRC','11','0','0','10','10','VEGF direction only (not stage-robust)'],
           ['STAD','10','0','0','9','8','3/5 direction only (ACRG)'],
           ['PAAD','10','0','0','0','0','none replicated (GSE21501)'],
           ['HCC','0','0','0','0','0','—'],
           ['ESCC','4','0','0','1','1','—'],
           ['IBD','6','6','11','1','0','—'],
           ['COPD','0','1','5','0','0','—'],
           ['NAFLD','0','0','0','0','0','—'],
           ['Asthma','0','0','4','0','0','—']])
        skip=True; continue
    if skip:
        if s.startswith(('## ','# ')): skip=False
        else: continue
    if s.startswith('# '): H(s[2:].strip(),0); continue
    if s.startswith('## '): H(s[3:].strip(),1); continue
    if s.startswith('### '): H(s[4:].strip(),2); continue
    doc.add_paragraph(s.strip())
fb=[('Figure 1. Analysis pipeline (v10 evidence layers).',[('','fig1_flowchart_v10.png')]),
    ('Figure 2. Module states. (A) raw; (B) joint composition models (diamonds = joint-robust; n.e. = not estimable).',[('A','fig2a_raw_v10.png'),('B','fig2b_joint_v10.png')]),
    ('Figure 3. Raw versus joint-model effects (170 states).',[('','fig3_composition_v10.png')]),
    ('Figure 4. TCGA module-OS associations (FDR<0.05).',[('','fig4_survival_v10.png')]),
    ('Figure 5. Molecular context. (A) amplification; (B) mutation.',[('','fig5_molecular_v10.png')]),
    ('Figure 6. Single-cell source hits (global FDR<0.05; patient numbers annotated).',[('','fig6_sc_v10.png')]),
    ('Supplementary Figure S1. Marker-proxy versus xCell medians.',[('','fig_paper_xcell_vs_marker.png')]),
    ('Supplementary Figure S2. xCell residualization heatmap (sensitivity).',[('','fig_paper_xcell_heatmap.png')])]
doc.add_page_break(); H('Figures',1)
for cap,pan in fb:
    H(cap,2)
    for tag,fn in pan:
        fp=os.path.join(P,fn)
        if tag: doc.add_paragraph(tag)
        if os.path.exists(fp): doc.add_picture(fp,width=Inches(6.2)); doc.add_paragraph()
        else: doc.add_paragraph('(missing: '+fn+')')
doc.save(os.path.join(V10,"01_稿件正文","Manuscript_draft_v10.docx"))
print("v10 docx written")
